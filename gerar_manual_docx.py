"""
Script para geração do Manual do Usuário em formato Microsoft Word (.docx).
Gera um documento profissional, estilizado e completo para o usuário final.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Define a cor de fundo de uma célula da tabela."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define o padding interno de uma célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="DICA IMPORTANTE"):
    """Cria uma caixa de destaque visual com borda lateral esquerda e fundo cinza claro."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F2F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E78"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}: ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(31, 78, 120)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(50, 50, 50)
    
    # Parágrafo vazio para espaçamento
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def formatar_tabela_padrao(table, col_widths, headers, data):
    """Aplica formatação elegante com cabeçalho azul e linhas intercaladas."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1F4E78")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = "F9FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(50, 50, 50)

def construir_manual(caminho_saida: Path):
    doc = docx.Document()
    
    # Margens da página: 2.0 cm
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # --- CABEÇALHO / TÍTULO PRINCIPAL ---
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(10)
    p_header.paragraph_format.space_after = Pt(4)
    run_sys = p_header.add_run("ROBÔ SOCIETÁRIO – DET / GOV.BR\n")
    run_sys.bold = True
    run_sys.font.name = "Calibri"
    run_sys.font.size = Pt(24)
    run_sys.font.color.rgb = RGBColor(31, 78, 120)
    
    run_sub = p_header.add_run("Manual de Instruções e Guia Prático Passo a Passo para o Usuário Final")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(89, 89, 89)
    
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("—" * 55)
    r_div.font.color.rgb = RGBColor(200, 200, 200)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # --- SEÇÃO 1: APRESENTAÇÃO E OBJETIVO ---
    h1 = doc.add_heading("1. Apresentação e Objetivo do Sistema", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 78, 120)
        
    p1 = doc.add_paragraph(
        "O Robô Societário é uma solução desenvolvida para automatizar de forma segura e ágil "
        "as consultas em lote no portal DET (Domicílio Eletrônico Trabalhista) do Ministério do Trabalho e Emprego (Gov.br). "
        "A ferramenta utiliza certificados digitais ICP-Brasil (A1 instalados no Windows) para realizar a autenticação "
        "e inspecionar as caixas postais das empresas indicadas."
    )
    p1.paragraph_format.space_after = Pt(6)
    
    doc.add_paragraph(
        "Dentre as principais funções automáticas da ferramenta, destacam-se:\n"
        " • Autenticação segura via Gov.br com seleção automatizada do certificado digital correspondente;\n"
        " • Troca dinâmica de perfil para procurador ou responsável legal conforme cadastrado;\n"
        " • Verificação de notificações e avisos pendentes de leitura ou já recebidos;\n"
        " • Download automático de termos de notificação e anexos oficiais em formato PDF ou ZIP;\n"
        " • Captura de telas (prints) em alta resolução de todas as etapas para auditoria e evidência jurídica;\n"
        " • Emissão de relatório consolidado em Excel contendo o status detalhado de cada CNPJ consultado."
    )
    
    # --- SEÇÃO 2: ESTRUTURA DOS ARQUIVOS ---
    h2 = doc.add_heading("2. Estrutura de Arquivos e Pastas do Sistema", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 78, 120)
        
    p2 = doc.add_paragraph(
        "O Robô Societário é portátil e autossuficiente: todos os arquivos, planilhas, relatórios e capturas "
        "ficam organizados na mesma pasta descompactada, sem a necessidade de instalar outros programas."
    )
    p2.paragraph_format.space_after = Pt(8)
    
    tabela_arquivos = doc.add_table(rows=1, cols=3)
    headers_arq = ["Item / Arquivo / Pasta", "Tipo", "Finalidade e Descrição"]
    widths_arq = [Inches(2.2), Inches(1.1), Inches(3.2)]
    dados_arq = [
        ["RoboSocietario.exe", "Aplicativo (.exe)", "Executável principal da automação. Basta dar 2 cliques para iniciar as consultas."],
        ["Atualizar_Certificados.bat", "Utilitário (.bat)", "Atalho rápido de 2 cliques para ler os certificados da máquina e preencher o Excel."],
        ["CNPJs.xlsx", "Planilha Excel", "Arquivo onde o usuário preenche a lista de CNPJs e indica o certificado correspondente."],
        ["CapturasDeTela-Downloads/", "Pasta de Saída", "Armazena todos os prints e evidências organizados por data, hora, certificado e CNPJ."],
        ["Relatorios/", "Pasta de Saída", "Armazena o relatório geral consolidado da execução em planilha Excel."],
        ["Logs/", "Pasta de Histórico", "Registra em arquivo .txt o histórico minuto a minuto da execução para auditoria e suporte."],
        ["CNPJs_Pendentes.xlsx", "Planilha Temporária", "Criada automaticamente apenas se algum CNPJ apresentar instabilidade no portal."]
    ]
    formatar_tabela_padrao(tabela_arquivos, widths_arq, headers_arq, dados_arq)
    
    p_spc = doc.add_paragraph()
    p_spc.paragraph_format.space_before = Pt(4)
    
    # --- SEÇÃO 3: GUIA PASSO A PASSO ---
    h3 = doc.add_heading("3. Guia Passo a Passo de Utilização", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 78, 120)
        
    # Passo 1
    h_p1 = doc.add_heading("Passo 1: Instalação dos Certificados Digitais na Máquina", level=2)
    for r in h_p1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "Certifique-se de que os certificados digitais A1 necessários estejam devidamente instalados "
        "no repositório pessoal do Windows da sua máquina (verificável pelo comando certmgr.msc -> Pessoal -> Certificados).\n"
        "O robô reconhece automaticamente todos os certificados instalados nesse repositório."
    )
    
    # Passo 2
    h_p2 = doc.add_heading("Passo 2: Atualização dos Certificados no Excel (2 Cliques)", level=2)
    for r in h_p2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "Sempre que você instalar um novo certificado na máquina ou for preparar uma nova consulta, "
        "basta dar dois cliques no arquivo utilitário:\n\n"
        "       👉 Atualizar_Certificados.bat\n\n"
        "Em aproximadamente 1 segundo, o script fará a varredura do Windows e atualizará a aba "
        "'Certificados_Disponiveis' da planilha CNPJs.xlsx com os nomes exatos, números de documento e datas de validade."
    )
    add_callout(
        doc,
        "Certifique-se de que a planilha CNPJs.xlsx não esteja aberta no Excel durante esse clique. "
        "O Windows impede alterações em arquivos que já estejam abertos por outro programa.",
        title="AVISO DE ARQUIVO FECHADO"
    )
    
    # Passo 3
    h_p3 = doc.add_heading("Passo 3: Preenchimento da Planilha CNPJs.xlsx", level=2)
    for r in h_p3.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "Abra o arquivo CNPJs.xlsx no Microsoft Excel. A planilha possui duas abas planejadas para facilitar o trabalho:\n\n"
        "1. Aba 'Certificados_Disponiveis' (Apenas Consulta):\n"
        "   Exibe todos os certificados instalados na sua máquina com Razão Social, CNPJ do Certificado, "
        "Validade e Status (VALIDO em verde ou EXPIRADO em vermelho).\n\n"
        "2. Aba 'CNPJs_consultados' (Preenchimento das Consultas):\n"
        "   É aqui onde você informa quais empresas serão consultadas pelo robô:"
    )
    
    tabela_exemplo = doc.add_table(rows=1, cols=3)
    headers_ex = ["Coluna", "Nome no Cabeçalho", "O que preencher (Exemplo)"]
    widths_ex = [Inches(1.2), Inches(2.0), Inches(3.3)]
    dados_ex = [
        ["Coluna A", "CNPJ", "Digite os 14 dígitos do CNPJ da empresa a ser consultada (ex: 39407892000185)."],
        ["Coluna B", "Certificado", "Digite a Razão Social ou apenas o CNPJ do certificado procurador (ex: 20791304000121)."]
    ]
    formatar_tabela_padrao(tabela_exemplo, widths_ex, headers_ex, dados_ex)
    
    p_spc2 = doc.add_paragraph()
    p_spc2.paragraph_format.space_before = Pt(4)
    
    add_callout(
        doc,
        "Você não precisa digitar o nome técnico complexo do certificado! "
        "Basta copiar apenas os 14 dígitos do CNPJ ou a Razão Social da aba 'Certificados_Disponiveis' "
        "e colar na coluna B. O robô possui correspondência inteligente e localiza o certificado correto automaticamente.",
        title="FACILIDADE DE PREENCHIMENTO"
    )
    
    doc.add_paragraph(
        "Importante: Após preencher e conferir a lista, salve o arquivo e feche o Excel antes de prosseguir."
    )
    
    # Passo 4
    h_p4 = doc.add_heading("Passo 4: Executando a Automação (RoboSocietario.exe)", level=2)
    for r in h_p4.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "Para iniciar as consultas, dê dois cliques diretamente no executável:\n\n"
        "       👉 RoboSocietario.exe\n\n"
        "O sistema abrirá uma janela preta (Console) e inicializará o Google Chrome automaticamente. "
        "A tela exibirá o progresso em tempo real, informando o item atual, a porcentagem de conclusão e a etapa em execução."
    )
    
    # Passo 5
    h_p5 = doc.add_heading("Passo 5: Tratamento Automático de Falhas e Pendências", level=2)
    for r in h_p5.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "O portal Gov.br e o DET podem apresentar lentidão ou quedas temporárias de conexão durante consultas em lote. "
        "Para garantir a integridade dos dados e evitar perda de trabalho, o Robô Societário possui um mecanismo inteligente e automatizado:\n\n"
        " • Salvamento Contínuo: Cada CNPJ consultado com sucesso tem seus prints, termos e dados salvos imediatamente no relatório geral e na pasta de evidências.\n"
        " • Segunda Passada Automática (Rodada 2): Caso algum CNPJ sofra instabilidade temporária no portal (timeout ou lentidão), o robô não interrompe a operação. Ele conclui a lista e inicia automaticamente uma segunda rodada para reprocessar os itens que falharam na primeira passada.\n"
        " • Planilha de Pendências (CNPJs_Pendentes.xlsx): Somente se algum CNPJ persistir com falha mesmo após a segunda tentativa ou exigir validação em duas etapas (2FA), o robô salva ao final da execução a planilha 'CNPJs_Pendentes.xlsx'.\n"
        " • Continuidade em Nova Execução: Havendo pendências registradas, ao abrir o RoboSocietario.exe o sistema permite escolher se deseja reprocessar apenas os pendentes restantes ou executar a base completa (CNPJs.xlsx).\n"
        " • Limpeza Automática: Quando todos os CNPJs forem concluídos com êxito, a planilha de pendências é excluída automaticamente."
    )
    
    # Passo 6
    h_p6 = doc.add_heading("Passo 6: Conferência dos Resultados e Evidências", level=2)
    for r in h_p6.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(46, 117, 182)
    doc.add_paragraph(
        "Ao finalizar a execução, o robô exibirá uma mensagem de conclusão e aguardará você pressionar ENTER. "
        "Você poderá conferir os resultados nos seguintes locais:\n\n"
        " 1. Pasta 'Relatorios':\n"
        "    Contém a planilha Relatorio_Execucao_<Data_Hora>.xlsx com o consolidado de todas as consultas, "
        "indicando se houve notificações, se havia anexos, quantidade de mensagens e links diretos para os prints.\n\n"
        " 2. Pasta 'CapturasDeTela-Downloads':\n"
        "    Organizada por data e hora da execução, separada por certificado e contendo uma subpasta dedicada para cada CNPJ com todos os "
        "prints comprobatórios e arquivos de anexos (PDFs/ZIPs) baixados."
    )
    
    # --- SEÇÃO 4: DICAS E PERGUNTAS FREQUENTES ---
    h4 = doc.add_heading("4. Recomendações e Perguntas Frequentes (FAQ)", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 78, 120)
        
    faq_items = [
        ("Posso utilizar o computador enquanto o robô está em execução?",
         "Durante a seleção do certificado digital na janela nativa do Windows, o robô precisa que essa janela "
         "esteja visível na tela. Após o login no portal, o restante da navegação e download ocorre no navegador. "
         "Recomenda-se evitar mexer no mouse ou teclado durante os primeiros segundos de troca de certificado."),
         
        ("Por que a janela do robô não fechou sozinha no final?",
         "A janela do console é mantida aberta intencionalmente com a mensagem 'Pressione ENTER para fechar...' "
         "para garantir que você consiga ler os totais processados e eventuais alertas com calma antes de encerrar."),
         
        ("O que fazer se um certificado venceu?",
         "Ao executar o 'Atualizar_Certificados.bat', verifique a coluna Status na aba 'Certificados_Disponiveis'. "
         "Certificados vencidos aparecem destacados em vermelho como 'EXPIRADO'. Providencie a renovação e "
         "instalação do novo certificado no Windows e clique no .bat novamente."),
         
        ("Posso consultar mais de 50 ou 100 CNPJs na mesma execução?",
         "Sim. O robô agrupa os CNPJs que utilizam o mesmo certificado para evitar logins desnecessários, "
         "otimizando o tempo total de processamento.")
    ]
    
    for perg, resp in faq_items:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(f"❓ {perg}")
        r_q.bold = True
        r_q.font.name = "Calibri"
        r_q.font.size = Pt(11)
        r_q.font.color.rgb = RGBColor(31, 78, 120)
        
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(resp)
        r_a.font.name = "Calibri"
        r_a.font.size = Pt(10)
        r_a.font.color.rgb = RGBColor(60, 60, 60)
        
    doc.save(str(caminho_saida))
    print(f"Documento gerado com sucesso em: {caminho_saida}")

if __name__ == "__main__":
    caminho = Path("Manual_do_Usuario_RoboSocietario.docx").resolve()
    construir_manual(caminho)
