import os
import re
import time
import shutil
import socket
import subprocess
import threading
import ctypes
from datetime import datetime
from pathlib import Path

import openpyxl
# pyrefly: ignore[missing-import]
from docx import Document
from pywinauto import Desktop
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

# ==============================================================================
# CONFIGURAÇÕES E CAMINHOS (Baseados no RoboSocietario-Principal-v10)
# ==============================================================================
BASE_DIR = Path(r"Z:\Temp\SOCIETARIO\AutomacaoCertificados")
LOG_PATH = BASE_DIR / "logs.docx"
EXCEL_PATH = BASE_DIR / "CNPJs-e-CPFs.xlsx"
SCREENSHOTS_DIR = BASE_DIR / "CapturasDeTela"

URL_LOGIN = "https://det.sit.trabalho.gov.br/login"
URL_SERVICOS = "https://det.sit.trabalho.gov.br/servicos"

# ModoTeste = True: considera notificações com data <= hoje (testes).
# Em produção, trocar para False (data >= hoje).
MODO_TESTE = True

# Configurações do Chrome Real via CDP (Evita CAPTCHA do gov.br)
CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
CHROME_USER_DATA = Path(os.environ["LOCALAPPDATA"]) / "Google" / "Chrome" / "User Data_RoboDET"
CDP_PORT = 9222

# Delay para acompanhamento visual do usuário (0.5s entre ações)
DELAY_ENTRE_ACOES_S = 0.5
DELAY_SLOW_MO_MS = 500


# ==============================================================================
# GERENCIADOR DE LOGS EM WORD (Substitui RoboSocietario-Log-v10)
# ==============================================================================
class WordLogger:
    def __init__(self, doc_path: Path):
        self.doc_path = doc_path
        self.doc_path.parent.mkdir(parents=True, exist_ok=True)
        if self.doc_path.exists():
            try:
                self.doc_path.unlink()
            except OSError:
                pass
        self.doc = Document()

    def _now_str(self) -> str:
        return datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    def log_etapa(self, item_id: str, etapa: str, msg: str):
        texto = f"{self._now_str()} - [{item_id}] [{etapa}] {msg}"
        print(texto)
        try:
            self.doc.add_paragraph(texto)
            self.doc.save(str(self.doc_path))
        except Exception as e:
            print(f"[AVISO] Falha ao salvar logs.docx: {e}")

    def log_raw(self, msg: str):
        texto = f"{self._now_str()} - {msg}"
        print(texto)
        try:
            self.doc.add_paragraph(texto)
            self.doc.save(str(self.doc_path))
        except Exception as e:
            print(f"[AVISO] Falha ao salvar logs.docx: {e}")


# ==============================================================================
# FUNÇÕES AUXILIARES (Navegador Real, Pop-up Nativo de Certificado e Pastas)
# ==============================================================================
def porta_em_uso(porta: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(("127.0.0.1", porta)) == 0


def iniciar_chrome_real():
    """
    Inicia o executável oficial do Google Chrome com perfil persistente e porta CDP.
    Como não é iniciado pelo binário de testes do Playwright, o gov.br não detecta
    flags de automação e mantém cookies/cache como no robô original.
    """
    if not porta_em_uso(CDP_PORT):
        CHROME_USER_DATA.mkdir(parents=True, exist_ok=True)
        cmd = [
            CHROME_PATH,
            f"--remote-debugging-port={CDP_PORT}",
            f"--user-data-dir={CHROME_USER_DATA}",
            "--start-maximized",
            "--no-first-run",
            "--no-default-browser-check",
            "--disable-blink-features=AutomationControlled",
        ]
        subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(20):
            if porta_em_uso(CDP_PORT):
                break
            time.sleep(0.5)


def fechar_navegador(browser=None, context=None, page=None):
    """
    Fecha a janela do Chrome real via comando CDP 'Browser.close'
    e encerra a sessão do Playwright de forma limpa.
    """
    try:
        if context and page and not page.is_closed():
            cdp = context.new_cdp_session(page)
            cdp.send("Browser.close")
            time.sleep(1)
    except Exception:
        pass

    try:
        if browser:
            browser.close()
    except Exception:
        pass


def confirmar_popup_certificado_nativo(timeout_segundos: int = 12):
    """
    Monitora em background a janela nativa do Windows/Navegador de seleção de
    Certificado Digital e clica em 'OK' (substitui o UIAutomation.Click do PAD).
    """
    fim = time.time() + timeout_segundos
    while time.time() < fim:
        try:
            desktop = Desktop(backend="uia")
            for win in desktop.windows():
                titulo = win.window_text()
                if any(k in titulo for k in ["gov.br", "Acesse sua conta", " Selecionar um certificado", " certificado"]):
                     btns_ok = win.descendants(control_type="Button", title="OK")
                     if btns_ok:
                         btns_ok[0].click_input()
                         return
        except Exception:
            pass
        time.sleep(0.5)


def limpar_ou_criar_pasta(caminho: Path):
    if not caminho.exists():
        caminho.mkdir(parents=True, exist_ok=True)
    else:
        for item in caminho.iterdir():
            try:
                if item.is_file() or item.is_symlink():
                    item.unlink()
                elif item.is_dir():
                    shutil.rmtree(item)
            except Exception as e:
                print(f"[AVISO] Não foi possível limpar {item}: {e}")


def remover_pastas_vazias_do_item(pasta_item: Path):
    """Replica a limpeza de subpastas vazias e da pasta principal do item."""
    for sub in ["CaixaPostal", "Notificacao", "ArquivosBaixados"]:
        sub_path = pasta_item / sub
        if sub_path.exists() and not any(sub_path.iterdir()):
            try:
                sub_path.rmdir()
            except OSError:
                pass

    if pasta_item.exists() and not any(pasta_item.iterdir()):
        try:
            pasta_item.rmdir()
        except OSError:
            pass


def exibir_alerta_planilha_inexistente():
    mensagem = (
        "A Planilha de referência para consulta de CNPJs e CPFs não existia.\n\n"
        f"A planilha foi criada na pasta CERTIFICADODET no caminho {BASE_DIR}\n"
        "Por favor, a preencha com os valores desejados, um abaixo do outro e execute o robô novamente"
    )
    ctypes.windll.user32.MessageBoxW(0, mensagem, "Planilha de Referência não existe", 0x30)


# ==============================================================================
# FLUXO PRINCIPAL DA AUTOMAÇÃO
# ==============================================================================
def executar_robo():
    logger = WordLogger(LOG_PATH)
    hoje_date = datetime.now().date()
    arquivos_pendentes_excluir = []

    # 1. Verificação da Planilha de Entrada
    if not EXCEL_PATH.exists():
        logger.log_raw("Planilha de CNPJs não encontrada.")
        wb = openpyxl.Workbook()
        wb.save(str(EXCEL_PATH))
        wb.close()
        exibir_alerta_planilha_inexistente()
        return

    # 2. Bloco de Preparação e Login
    playwright_mgr = None
    browser = None
    context = None
    page = None
    try:
        wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
        sheet = wb.active
        excel_data = [
            str(cell.value).strip()
            for row in sheet.iter_rows(values_only=False)
            for cell in row
            if cell.value is not None and str(cell.value).strip() != ""
        ]
        wb.close()

        limpar_ou_criar_pasta(SCREENSHOTS_DIR)

        iniciar_chrome_real()
        playwright_mgr = sync_playwright().start()
        browser = playwright_mgr.chromium.connect_over_cdp(
            f"http://127.0.0.1:{CDP_PORT}",
            slow_mo=DELAY_SLOW_MO_MS
        )
        context = browser.contexts[0]
        page = context.pages[0] if context.pages else context.new_page()

        page.goto(URL_LOGIN, timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)

        # Verifica até 3 vezes se o botão "Entrar com gov.br" carregou
        btn_govbr = page.locator("button#botao, button:has-text('Entrar com gov.br')").first
        pagina_carregou = False
        for _ in range(3):
            if btn_govbr.is_visible():
                pagina_carregou = True
                break
            time.sleep(2)

        if not pagina_carregou:
            logger.log_raw(
                "Não foi possível abrir o navegador e identificar os valores necessários para o robô continuar."
            )
            fechar_navegador(browser, context, page)
            if playwright_mgr:
                playwright_mgr.stop()
            return

        # Clica em "Entrar com gov.br"
        time.sleep(DELAY_ENTRE_ACOES_S)
        btn_govbr.click()
        time.sleep(DELAY_ENTRE_ACOES_S)

        # Aguarda botão "Seu certificado digital" e dispara thread para clicar no "OK" nativo
        btn_cert = page.locator("button:has-text('Seu certificado digital'), #login-certificate").first
        btn_cert.wait_for(state="visible", timeout=60000)

        thread_popup = threading.Thread(target=confirmar_popup_certificado_nativo, args=(15,), daemon=True)
        thread_popup.start()
        time.sleep(DELAY_ENTRE_ACOES_S)
        btn_cert.click(no_wait_after=True)

        # Aguarda carregar a tela de Serviços (Caixa Postal)
        card_caixa_postal = page.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))")
        card_caixa_postal.wait_for(state="visible", timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)

    except Exception as e:
        # Equivalente ao RoboSocietario-Erro-v10.txt
        logger.log_raw(f"Erro fatal no robô (fase de preparação ou login): {e}")
        fechar_navegador(browser, context, page)
        if playwright_mgr:
            playwright_mgr.stop()
        raise

    # 3. Loop de Processamento de CPFs / CNPJs
    perfil = 0
    falhas_seguidas = 0
    etapa = "Início"

    for current_item in excel_data:
        if falhas_seguidas >= 3:
            logger.log_raw("Execução interrompida: 3 falhas seguidas. Os itens restantes não foram processados.")
            break

        item_id = current_item.strip()
        if not item_id:
            continue

        try:
            etapa = "Lendo o item da planilha"
            logger.log_etapa(item_id, etapa, "Iniciando o processamento do item")

            etapa = "Aguardando a tela de serviços"
            page.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))").wait_for(state="visible", timeout=60000)
            time.sleep(DELAY_ENTRE_ACOES_S)

            etapa = "Trocando de perfil"
            btn_trocar_perfil = page.locator("button:has-text('Trocar Perfil')").first
            btn_trocar_perfil.wait_for(state="visible", timeout=20000)
            time.sleep(DELAY_ENTRE_ACOES_S)
            btn_trocar_perfil.click()
            time.sleep(2)

            # Clica no dropdown de seleção de perfil e navega via teclado como no PAD
            dropdown_perfil = page.locator("ng-select, .br-select, span.ng-arrow-wrapper").first
            dropdown_perfil.click()
            time.sleep(DELAY_ENTRE_ACOES_S)

            if perfil == 0:
                perfil += 1
                page.keyboard.press("ArrowDown")
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.keyboard.press("Enter")
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.keyboard.press("Tab")
            else:
                page.keyboard.press("ArrowDown")
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.keyboard.press("ArrowDown")
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.keyboard.press("Enter")
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.keyboard.press("Tab")

            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.type(item_id, delay=100)
            time.sleep(DELAY_ENTRE_ACOES_S)

            etapa = "Selecionando o perfil e validando CPF/CNPJ"
            page.locator("button:has-text('Selecionar')").first.click()
            time.sleep(2)

            # Verifica se apareceu mensagem de CPF/CNPJ inválido no DOM (substitui OCR)
            msg_invalido = page.get_by_text(re.compile(r"(CNPJ|CPF)\s+inv.lido", re.IGNORECASE))
            if msg_invalido.count() > 0 and msg_invalido.first.is_visible():
                time.sleep(DELAY_ENTRE_ACOES_S)
                page.locator("button:has-text('Cancelar')").first.click()
                time.sleep(DELAY_ENTRE_ACOES_S)
                logger.log_raw(f"Documento inválido, item ignorado: {item_id}")
                continue

            # Criação de pastas do item
            etapa = "Criando a pasta do item"
            pasta_item = SCREENSHOTS_DIR / item_id
            pasta_caixa_postal = pasta_item / "CaixaPostal"
            pasta_notificacao = pasta_item / "Notificacao"
            pasta_arquivos_baixados = pasta_item / "ArquivosBaixados"

            etapa = "Criando a subpasta CaixaPostal"
            pasta_caixa_postal.mkdir(parents=True, exist_ok=True)
            etapa = "Criando a subpasta Notificacao"
            pasta_notificacao.mkdir(parents=True, exist_ok=True)
            etapa = "Criando a subpasta ArquivosBaixados"
            pasta_arquivos_baixados.mkdir(parents=True, exist_ok=True)
            time.sleep(DELAY_ENTRE_ACOES_S)

            # --- CAIXA POSTAL ---
            etapa = "Abrindo a Caixa Postal"
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))").click()
            time.sleep(DELAY_ENTRE_ACOES_S)
            formatted_dt = datetime.now().strftime("%d%m%Y_%H%M%S")

            page.locator("span:has-text('Exibir apenas não lidas')").first.click()
            time.sleep(DELAY_ENTRE_ACOES_S)

            etapa = "Tirando o print da Caixa Postal"
            # Verifica se há itens listados na Caixa Postal (aguarda até 10s sem quebrar se vazio)
            linhas_cp = page.locator("ngx-datatable datatable-body-row, table tbody tr")
            try:
                linhas_cp.first.wait_for(state="visible", timeout=10000)
                tem_msg_nao_lida = True
            except PlaywrightTimeoutError:
                tem_msg_nao_lida = False

            if tem_msg_nao_lida:
                time.sleep(DELAY_ENTRE_ACOES_S)
                caminho_print_cp = pasta_caixa_postal / f"{item_id}_{formatted_dt}.jpg"
                page.screenshot(path=str(caminho_print_cp), type="jpeg")
                logger.log_etapa(item_id, etapa, "Print da Caixa Postal realizado: existe mensagem não lida.")
                time.sleep(DELAY_ENTRE_ACOES_S)

            # Volta para a tela de Serviços
            time.sleep(DELAY_ENTRE_ACOES_S)
            btn_voltar = page.locator("a:has-text('Serviços'), button:has-text('Voltar'), i.fa-arrow-left, i.fa-home").first
            if btn_voltar.is_visible():
                btn_voltar.click()
            else:
                page.goto(URL_SERVICOS, timeout=60000)
            time.sleep(DELAY_ENTRE_ACOES_S)

            # --- NOTIFICAÇÃO ---
            etapa = "Abrindo a tela de Notificação"
            card_notificacao = page.locator(".cardListItem:has(span.title:has-text('NOTIFICAÇÃO'))")
            card_notificacao.wait_for(state="visible", timeout=60000)
            time.sleep(DELAY_ENTRE_ACOES_S)
            card_notificacao.click()
            time.sleep(DELAY_ENTRE_ACOES_S)

            etapa = "Tirando o print da Notificação"
            linhas_notif = page.locator("ngx-datatable datatable-row-wrapper")
            try:
                linhas_notif.first.wait_for(state="visible", timeout=10000)
                tem_notificacoes = True
            except PlaywrightTimeoutError:
                tem_notificacoes = False

            etapa = "Lendo a lista de notificações"
            if tem_notificacoes:
                qtd_linhas = linhas_notif.count()
                logger.log_etapa(
                    item_id,
                    etapa,
                    f"Lista de notificações lida: {qtd_linhas} linha(s) extraída(s). Data de hoje considerada: {hoje_date.strftime('%d/%m/%Y')}",
                )

                # Extrai os dados das linhas (Célula 2: valor 1 | Célula 5: Data dd/MM/yyyy)
                dados_linhas = []
                existe_notificacao_valida = False

                for idx in range(qtd_linhas):
                    row_loc = linhas_notif.nth(idx)
                    texto_data = row_loc.locator("datatable-body-cell").nth(5).inner_text().strip()
                    try:
                        data_notif = datetime.strptime(texto_data, "%d/%m/%Y").date()
                        valida = (data_notif <= hoje_date) if MODO_TESTE else (data_notif >= hoje_date)
                    except ValueError:
                        data_notif = None
                        valida = False

                    if valida:
                        existe_notificacao_valida = True
                    dados_linhas.append((idx, texto_data, valida))

                if not existe_notificacao_valida:
                    logger.log_etapa(item_id, etapa, "Nenhuma notificação possui data válida e o print foi ignorado.")

                # Processa cada linha da lista de notificações
                for idx, texto_data, linha_valida in dados_linhas:
                    etapa = f"Lendo a data da linha {idx}"
                    if linha_valida:
                        row_loc = linhas_notif.nth(idx)

                        # Rola a tela até a linha antes de tirar o print recortado
                        etapa = f"Rolando a página até a linha {idx}"
                        row_loc.scroll_into_view_if_needed()
                        time.sleep(DELAY_ENTRE_ACOES_S)

                        # Print recortado exclusivamente no retângulo da linha da notificação
                        etapa = f"Tirando print recortado da linha {idx}"
                        caminho_print_linha = pasta_notificacao / f"{item_id}_linha_{idx}_{formatted_dt}.jpg"
                        row_loc.screenshot(path=str(caminho_print_linha), type="jpeg")
                        time.sleep(DELAY_ENTRE_ACOES_S)
                        logger.log_etapa(
                            item_id,
                            etapa,
                            f"Print recortado da linha {idx} (data {texto_data}) realizado com sucesso.",
                        )

                        # Localiza e realiza o download caso a linha possua anexo disponível
                        etapa = f"Localizando o botão de download da linha {idx}"
                        btn_download = row_loc.locator("datatable-body-cell").nth(8).locator("button").nth(1)
                        if btn_download.count() == 0:
                            btn_download = row_loc.locator("button:has(i.fa-download), button[title*='Download'], button[title*='Baixar']").first

                        if btn_download.count() > 0 and btn_download.is_visible():
                            logger.log_etapa(item_id, etapa, f"Linha {idx} (data {texto_data}): clicando no botão de download")
                            try:
                                etapa = f"Aguardando o download da linha {idx}"
                                time.sleep(DELAY_ENTRE_ACOES_S)
                                with page.expect_download(timeout=60000) as download_info:
                                    btn_download.click()
                                    logger.log_etapa(item_id, etapa, f"Linha {idx}: clique no botão de download realizado")

                                download = download_info.value
                                destino_final = pasta_arquivos_baixados / download.suggested_filename
                                download.save_as(str(destino_final))
                                time.sleep(DELAY_ENTRE_ACOES_S)
                                logger.log_etapa(
                                    item_id,
                                    etapa,
                                    f"Linha {idx}: arquivo baixado e movido para ArquivosBaixados ({destino_final})",
                                )
                            except PlaywrightTimeoutError:
                                logger.log_etapa(
                                    item_id,
                                    etapa,
                                    f"AVISO: download não detectado em 60s (linha {idx})",
                                )
                        else:
                            logger.log_etapa(
                                item_id,
                                etapa,
                                f"Linha {idx} (data {texto_data}): notificação válida sem botão de download (apenas leitura).",
                            )
                    else:
                        logger.log_etapa(
                            item_id,
                            etapa,
                            f"Linha {idx} ignorada: data {texto_data} fora do critério (ModoTeste={MODO_TESTE}, hoje={hoje_date.strftime('%d/%m/%Y')})",
                        )
            else:
                logger.log_etapa(item_id, etapa, "Nenhuma notificação listada na tela de Notificação")

            # Retorna à tela de serviços
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.goto(URL_SERVICOS, timeout=60000)
            time.sleep(DELAY_ENTRE_ACOES_S)

            # Limpeza de diretórios vazios e contabilização
            etapa = "Limpando subpastas vazias"
            remover_pastas_vazias_do_item(pasta_item)

            etapa = "Conferindo os arquivos e limpando a pasta vazia"
            arquivos_guardados = list(pasta_item.rglob("*")) if pasta_item.exists() else []
            qtd_arquivos = sum(1 for f in arquivos_guardados if f.is_file())

            logger.log_raw(
                f"SUCESSO: O item '{item_id}' foi verificado. Arquivos guardados (prints + downloads): {qtd_arquivos}"
            )
            falhas_seguidas = 0

        except Exception as e:
            # Equivalente ao RoboSocietario-Erro_cpf-v10.txt
            falhas_seguidas += 1
            logger.log_raw(f"Erro no item '{current_item}' na etapa '{etapa}': {e}")
            try:
                page.goto(URL_SERVICOS, timeout=60000)
            except Exception:
                pass

    # 4. Encerramento e Logout
    try:
        time.sleep(DELAY_ENTRE_ACOES_S)
        logger.log_raw("Finalizando execução. Efetuando logout no sistema...")

        # No DET o botão é um br-button ou button com o texto 'Sair' no cabeçalho
        btn_sair = page.locator("br-button:has-text('Sair'), button:has-text('Sair'), a:has-text('Sair')").first
        if not btn_sair.is_visible():
            btn_sair = page.get_by_text(re.compile(r"^\s*Sair\s*$", re.IGNORECASE)).first

        if btn_sair.is_visible():
            time.sleep(DELAY_ENTRE_ACOES_S)
            btn_sair.click()
            logger.log_raw("Clique no botão 'Sair' realizado.")
            time.sleep(DELAY_ENTRE_ACOES_S)

            # Caso apareça popup ou modal de confirmação (ex: Sim / Confirmar)
            btn_confirmar = page.locator("button:has-text('Sim'), button:has-text('Confirmar'), br-button:has-text('Sim')").first
            try:
                if btn_confirmar.is_visible(timeout=2000):
                    btn_confirmar.click()
                    time.sleep(DELAY_ENTRE_ACOES_S)
            except Exception:
                pass

            # Aguarda a requisição de logout processar
            time.sleep(1.5)
        else:
            logger.log_raw("Botão 'Sair' não localizado no cabeçalho.")
    except Exception as e:
        logger.log_raw(f"Aviso durante tentativa de logout: {e}")

    for arq in arquivos_pendentes_excluir:
        try:
            Path(arq).unlink(missing_ok=True)
        except Exception:
            pass

    logger.log_raw("Fechando o navegador Chrome...")
    fechar_navegador(browser, context, page)
    if playwright_mgr:
        playwright_mgr.stop()
    logger.log_raw("Automação finalizada com sucesso.")


if __name__ == "__main__":
    executar_robo()