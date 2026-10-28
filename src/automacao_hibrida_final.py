import os
import sys
import warnings

# Define STA (Single-Threaded Apartment) antes de qualquer importação do pywinauto/comtypes
if not hasattr(sys, "coinit_flags"):
    sys.coinit_flags = 2
warnings.filterwarnings("ignore", message=".*Apply externally defined coinit_flags.*")
import re
import time
import shutil
import socket
import subprocess
import ctypes
import logging
import threading
from typing import Optional, Callable
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from PIL import Image

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pyautogui
pyautogui.FAILSAFE = False
from pywinauto import Desktop
from pywinauto.keyboard import send_keys
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

# ==============================================================================
# CONFIGURAÇÕES E CAMINHOS
# ==============================================================================
URL_LOGIN = "https://det.sit.trabalho.gov.br/login"
URL_SERVICOS = "https://det.sit.trabalho.gov.br/servicos"

# ModoTeste = True: considera notificações com data <= hoje (testes).
# Em produção, trocar para False (data >= hoje).
MODO_TESTE = False

# Diretório base padrão para depuração e caminhos de fallback
BASE_DIR = Path.cwd()

# Configurações do Chrome Real e Porta CDP
CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
CHROME_USER_DATA = Path(os.environ["LOCALAPPDATA"]) / "Google" / "Chrome" / "User Data_RoboSocietario"
CDP_PORT = 9222

# Delays para acompanhamento visual do usuário (0.5s entre ações)
DELAY_ENTRE_ACOES_S = 0.5
DELAY_SLOW_MO_MS = 500

def obter_caminho_recurso(caminho_relativo: str) -> Path:
    """
    Retorna o caminho absoluto do recurso, compatível com execução direta (Python),
    PyInstaller em modo diretório (_internal / pasta do executável) ou modo temporário (_MEIPASS).
    """
    if getattr(sys, "frozen", False):
        base_meipass = getattr(sys, "_MEIPASS", None)
        if base_meipass:
            c = Path(base_meipass) / caminho_relativo
            if c.exists():
                return c

        exe_dir = Path(sys.executable).resolve().parent
        c_exe = exe_dir / caminho_relativo
        if c_exe.exists():
            return c_exe

        c_internal = exe_dir / "_internal" / caminho_relativo
        if c_internal.exists():
            return c_internal

    script_dir = Path(__file__).resolve().parent
    if (script_dir / caminho_relativo).exists():
        return script_dir / caminho_relativo
    if (script_dir.parent / caminho_relativo).exists():
        return script_dir.parent / caminho_relativo

    return script_dir.parent / caminho_relativo


def sanitizar_nome_pasta(cert_nome: str) -> str:
    """
    Sanitiza o nome do certificado para uso seguro como nome de pasta no Windows,
    removendo caracteres inválidos (\\/*?:"<>|) e limitando o tamanho.
    Exemplo: 'FALAVINHA NEXT ASSESSORIA:20791304000121' -> 'FALAVINHA NEXT ASSESSORIA_20791304000121'
    """
    if not cert_nome:
        return "Certificado_Padrao"
    nome_limpo = re.sub(r'[\\/*?:"<>|]', "_", str(cert_nome).strip())
    nome_limpo = re.sub(r'\s+', " ", nome_limpo)
    nome_limpo = re.sub(r'_+', "_", nome_limpo).strip(" ._")
    if not nome_limpo:
        return "Certificado"
    return nome_limpo[:60].strip(" ._")


# Imagens de referência para o login físico via PyAutoGUI (procura em assets/ e fallback)
_caminho_assets = obter_caminho_recurso("assets/capturasDeTela")
if _caminho_assets.exists():
    IMG_DIR = _caminho_assets
else:
    IMG_DIR = obter_caminho_recurso("capturasDeTela")

IMG_GOVBR = IMG_DIR / "EntrarComGov.png"
IMG_CERT = IMG_DIR / "SeuCertificadoDigital.png"
IMG_OK = IMG_DIR / "Ok.png"
IMG_CAPTCHA_AVANCAR = IMG_DIR / "CaptchaAvancar.png"
IMG_CAPTCHA_HAND = IMG_DIR / "CaptchaHand.png"
IMG_DEBUG_ERRO = IMG_DIR / "debug_tela_erro.png"


class CaptchaDetectadoException(Exception):
    """Exceção levantada quando um desafio hCaptcha é detectado durante o login gov.br."""
    pass


class VerificacaoDuasEtapasException(Exception):
    """Exceção levantada quando a tela de Verificação em duas etapas (2FA) do gov.br é detectada."""
    pass

# ==============================================================================
# LOGGER DA APLICAÇÃO (RoboSocietario)
# ==============================================================================
logger = logging.getLogger("RoboSocietario")


def log_etapa(item_id: str, etapa: str, msg: str):
    logger.info(f"[{item_id}] [{etapa}] {msg}")


# ==============================================================================
# FUNÇÕES DE SISTEMA OPERACIONAL (PORTAS, CHROME E PYAUTOGUI)
# ==============================================================================
def porta_em_uso(porta: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(("127.0.0.1", porta)) == 0


def matar_processos_chrome_automacao():
    """
    Encerra de forma forçada e definitiva todos os processos do Chrome vinculados
    à automação (porta 9222 ou pasta User Data_RoboSocietario), matando a árvore
    completa de processos para evitar processos zumbis sem afetar o Chrome pessoal do usuário.
    """
    global CHROME_PROC

    # 1. Se temos a referência do subprocess.Popen, encerra a árvore inteira via taskkill /F /T
    if CHROME_PROC is not None:
        try:
            pid = CHROME_PROC.pid
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(pid)],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=5,
            )
        except Exception:
            pass
        CHROME_PROC = None

    # 2. Encerra qualquer processo escutando na porta CDP (9222)
    try:
        resultado = subprocess.run(
            ["netstat", "-ano", "-p", "tcp"],
            capture_output=True,
            text=True,
            timeout=5,
        )
        pids_para_matar = set()
        for linha in resultado.stdout.splitlines():
            partes = linha.strip().split()
            if len(partes) >= 5 and f":{CDP_PORT}" in partes[1] and "LISTENING" in partes:
                pid_str = partes[-1]
                if pid_str.isdigit() and int(pid_str) > 0:
                    pids_para_matar.add(pid_str)
        for p_id in pids_para_matar:
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", p_id],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=5,
            )
    except Exception:
        pass

    # 3. Como garantia adicional, encerra processos chrome.exe associados ao perfil 'User Data_RoboSocietario'
    try:
        ps_cmd = (
            "& { Get-CimInstance Win32_Process | "
            "Where-Object { $_.Name -eq 'chrome.exe' -and $_.CommandLine -like '*User Data_RoboSocietario*' } | "
            "ForEach-Object { Stop-Process -Id $_.ProcessId -Force } }"
        )
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps_cmd],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=6,
        )
    except Exception:
        pass


def liberar_porta_se_ocupada():
    """Fecha processos órfãos na porta CDP para garantir uma inicialização limpa."""
    if porta_em_uso(CDP_PORT):
        logger.info(f"Porta {CDP_PORT} em uso. Encerrando processos anteriores da automação...")
        matar_processos_chrome_automacao()
        time.sleep(1)


CHROME_PROC: Optional[subprocess.Popen] = None


def iniciar_chrome_hibrido():
    """
    Inicia o executável oficial do Google Chrome com perfil persistente e porta CDP.
    O Playwright NÃO é conectado durante a inicialização, evitando qualquer flag ou
    sinalização de automação no gov.br.
    """
    global CHROME_PROC
    CHROME_USER_DATA.mkdir(parents=True, exist_ok=True)
    cmd = [
        CHROME_PATH,
        f"--remote-debugging-port={CDP_PORT}",
        f"--user-data-dir={CHROME_USER_DATA}",
        "--start-maximized",
        "--no-first-run",
        "--no-default-browser-check",
        "--disable-blink-features=AutomationControlled",
        URL_LOGIN,
    ]
    CHROME_PROC = subprocess.Popen(cmd)
    for _ in range(20):
        if porta_em_uso(CDP_PORT):
            break
        time.sleep(0.5)


def focar_janela_segura(win, *args, **kwargs):
    """
    Garante foco e primeiro plano na janela alvo de forma limpa e confiável,
    superando as restrições de Foreground Lock / Focus-Stealing do Windows
    (mesmo em execuções de 2ª rodada ou via Agendador de Tarefas), utilizando
    SwitchToThisWindow e permissões de thread cruzadas.
    """
    try:
        hwnd = win.handle if hasattr(win, "handle") else int(win)
        user32 = ctypes.windll.user32
        kernel32 = ctypes.windll.kernel32

        # 1. Se estiver minimizada, restaura
        if user32.IsIconic(hwnd):
            user32.ShowWindow(hwnd, 9)  # SW_RESTORE

        # 2. Conecta as threads do Windows (primeiro plano, janela alvo e processo da automação)
        h_fore = user32.GetForegroundWindow()
        target_thread = user32.GetWindowThreadProcessId(hwnd, None)
        fore_thread = user32.GetWindowThreadProcessId(h_fore, None) if h_fore else 0
        app_thread = kernel32.GetCurrentThreadId()

        attached_fore = False
        attached_app = False

        if fore_thread and target_thread and fore_thread != target_thread:
            attached_fore = bool(user32.AttachThreadInput(fore_thread, target_thread, True))
        if app_thread and target_thread and app_thread != target_thread:
            attached_app = bool(user32.AttachThreadInput(app_thread, target_thread, True))

        try:
            # Pulso rápido de ALT para resetar o bloqueio de roubo de foco (ForegroundLockTimeout)
            user32.keybd_event(0x12, 0, 0, 0)
            user32.keybd_event(0x12, 0, 2, 0)

            # SwitchToThisWindow com fAltTab=True simula o Alt-Tab nativo diretamente para a janela alvo
            user32.SwitchToThisWindow(hwnd, True)
            user32.BringWindowToTop(hwnd)
            user32.SetForegroundWindow(hwnd)
        finally:
            if attached_fore:
                user32.AttachThreadInput(fore_thread, target_thread, False)
            if attached_app:
                user32.AttachThreadInput(app_thread, target_thread, False)

        time.sleep(0.5)
    except Exception:
        pass


def obter_janela_chrome(timeout_segundos: int = 30):
    """Localiza exclusivamente a janela principal do Google Chrome aberta."""
    desktop = Desktop(backend="uia")
    fim = time.time() + timeout_segundos
    while time.time() < fim:
        try:
            for win in desktop.windows():
                txt = win.window_text()
                cls = win.class_name()
                txt_lower = txt.lower()
                # Ignora expressamente Antigravity, VS Code, Edge e Agendador de Tarefas
                if any(ignorar in txt_lower for ignorar in [
                    "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                    "agendador de tarefas", "task scheduler"
                ]):
                    continue
                # Identifica exclusivamente a janela do Google Chrome oficial
                if cls == "Chrome_WidgetWin_1" and ("google chrome" in txt_lower or txt_lower.endswith("chrome")):
                    focar_janela_segura(win)
                    return win
        except Exception:
            pass
        time.sleep(0.8)
    raise TimeoutError("Janela do Google Chrome não encontrada dentro do tempo limite.")


def clicar_robusto(pos_x: int, pos_y: int):
    """Move o mouse e dispara um clique físico completo e deliberado (mouseDown + mouseUp)."""
    pyautogui.moveTo(pos_x, pos_y, duration=0.6)
    time.sleep(0.25)
    pyautogui.mouseDown()
    time.sleep(0.2)
    pyautogui.mouseUp()
    time.sleep(0.2)


def localizar_imagem(caminho_imagem: Path, confidence: float = 0.8):
    """Localiza o centro da imagem na tela com tolerância a ClearType (grayscale)."""
    if not caminho_imagem.exists():
        return None

    try:
        pos = pyautogui.locateCenterOnScreen(str(caminho_imagem), confidence=confidence, grayscale=True)
        if pos:
            return pos
    except Exception:
        pass

    try:
        pos = pyautogui.locateCenterOnScreen(str(caminho_imagem), confidence=max(0.65, confidence - 0.1), grayscale=True)
        if pos:
            return pos
    except Exception:
        pass

    return None


def detectar_presenca_captcha() -> bool:
    """
    Verifica se o desafio hCaptcha surgiu na tela após clicar em 'Seu certificado digital'
    (ou durante a autenticação gov.br).
    Utiliza 3 métodos complementares com validação rigorosa para evitar falsos positivos:
    1. Detecção cromática por contagem de pixels na cor turquesa/teal (#028490) do hCaptcha.
    2. Correspondência de imagem PyAutoGUI (ícone da mãozinha ou botão 'Avançar' com validação cruzada).
    3. Inspeção UIA no Windows procurando por termos explícitos do desafio hCaptcha.
    """
    teal_count = 0
    # 1. Detecção cromática via screenshot (instantânea e imune a variações de texto)
    try:
        shot = pyautogui.screenshot()
        w, h = shot.size
        # Foca na região central da tela onde o card flutuante do hCaptcha é renderizado
        box = (int(w * 0.15), int(h * 0.05), int(w * 0.85), int(h * 0.95))
        cropped = shot.crop(box)
        for p in cropped.get_flattened_data():
            # R < 40, 95 < G < 155, 110 < B < 170 (cor sólida da barra superior do hCaptcha)
            if p[0] < 40 and 95 < p[1] < 155 and 110 < p[2] < 170:
                teal_count += 1
                if teal_count >= 1500:
                    logger.warning(
                        f"[CAPTCHA] Desafio hCaptcha detectado por padrão cromático ({teal_count} pixels turquesa identificados)!"
                    )
                    return True
    except Exception as e_cor:
        logger.debug(f"Erro na checagem cromática de captcha: {e_cor}")

    # 2. Correspondência de imagem (ícone de acessibilidade do hCaptcha ou botão Avançar validado)
    try:
        # Ícone de acessibilidade exclusivo do widget hCaptcha (mãozinha)
        if IMG_CAPTCHA_HAND.exists():
            pos_hand = pyautogui.locateCenterOnScreen(str(IMG_CAPTCHA_HAND), confidence=0.85)
            if pos_hand:
                logger.warning(f"[CAPTCHA] Desafio hCaptcha detectado pelo ícone da mãozinha em ({pos_hand.x}, {pos_hand.y})!")
                return True

        # O botão 'Avançar' só confirma captcha se tiver alta fidelidade E comprovação visual (cor turquesa)
        if IMG_CAPTCHA_AVANCAR.exists():
            pos_avancar = pyautogui.locateCenterOnScreen(str(IMG_CAPTCHA_AVANCAR), confidence=0.85)
            if pos_avancar:
                if teal_count >= 300:
                    logger.warning(
                        f"[CAPTCHA] Desafio hCaptcha detectado pelo botão 'Avançar' em ({pos_avancar.x}, {pos_avancar.y}) "
                        f"confirmado por {teal_count} pixels turquesa!"
                    )
                    return True
                else:
                    logger.debug(
                        f"Botão semelhante a 'Avançar' em ({pos_avancar.x}, {pos_avancar.y}) ignorado por falta de indícios cromáticos do hCaptcha."
                    )
    except Exception as e_img:
        logger.debug(f"Erro na checagem por imagem de captcha: {e_img}")

    # 3. Fallback UIA no Windows procurando por termos inequívocos do desafio
    try:
        desktop = Desktop(backend="uia")
        for win in desktop.windows():
            titulo = win.window_text().lower()
            if any(ignorar in titulo for ignorar in [
                "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                "agendador de tarefas", "task scheduler"
            ]):
                continue
            if any(k in titulo for k in ["gov.br", "acesse sua conta"]):
                for desc in win.descendants():
                    txt = (desc.window_text() or desc.element_info.name or "").lower()
                    if any(term in txt for term in ["clique nos itens", "clique em cada", "hcaptcha", "desafio"]):
                        logger.warning(f"[CAPTCHA] Desafio hCaptcha detectado via UIA pelo texto: '{txt}'!")
                        return True
    except Exception:
        pass

    return False


def detectar_presenca_2fa_uia() -> bool:
    """
    Verifica via UIA no Windows se a tela de Verificação em duas etapas (2FA) do gov.br está visível.
    """
    try:
        desktop = Desktop(backend="uia")
        for win in desktop.windows():
            titulo = win.window_text().lower()
            if any(ignorar in titulo for ignorar in [
                "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                "agendador de tarefas", "task scheduler"
            ]):
                continue
            if any(k in titulo for k in ["gov.br", "acesse sua conta", "chrome"]):
                for desc in win.children():
                    txt = (desc.window_text() or desc.element_info.name or "").lower()
                    if any(term in txt for term in [
                        "verificação em duas etapas", "verificacao em duas etapas",
                        "duas etapas", "código de acesso", "codigo de acesso",
                        "gerar código de acesso", "gerar codigo de acesso", "confirme sua identidade"
                    ]):
                        logger.warning(f"[2FA] Tela de Verificação em duas etapas detectada via UIA: '{txt}'!")
                        return True
    except Exception:
        pass
    return False


def detectar_verificacao_duas_etapas_playwright(page) -> bool:
    """
    Verifica se a página atual do gov.br está exibindo a tela de 'Verificação em duas etapas' (2FA).
    Detecta via avaliação direta de texto no DOM (instantâneo), seletores e HTML raw.
    """
    try:
        if page.is_closed():
            return False

        url_atual = ""
        try:
            url_atual = page.url.lower()
        except Exception:
            pass

        # 1. Avaliação rápida de texto visível no DOM via JavaScript (em milissegundos)
        try:
            texto_dom = page.evaluate("() => document.body ? document.body.innerText.toLowerCase() : ''")
        except Exception:
            texto_dom = ""

        termos_chave_2fa = [
            "confirme sua identidade com a verificação em duas etapas",
            "confirme sua identidade com a verificacao em duas etapas",
            "verificação em duas etapas",
            "verificacao em duas etapas",
            "gerar código de acesso",
            "gerar codigo de acesso",
            "não solicitar verificação em duas etapas",
            "nao solicitar verificacao em duas etapas",
        ]

        if any(t in texto_dom for t in termos_chave_2fa):
            logger.warning("[2FA] Tela de Verificação em duas etapas detectada via texto do DOM no Playwright!")
            return True

        if ("código de acesso" in texto_dom or "codigo de acesso" in texto_dom) and \
           any(k in url_atual for k in ["sso.acesso.gov.br", "acesso.gov.br", "login", "duasetapas"]):
            logger.warning("[2FA] Campo de código de acesso em URL gov.br detectado!")
            return True

        # 2. Seletor Playwright direto se estiver em URL de autenticação
        eh_url_auth = any(k in url_atual for k in ["sso.acesso.gov.br", "acesso.gov.br", "duasetapas", "login"])
        if eh_url_auth:
            seletores_2fa = [
                "h1:has-text('Confirme sua identidade')",
                "text=Confirme sua identidade com a Verificação em duas etapas",
                "text=Verificação em duas etapas",
                "input[placeholder*='código de acesso' i]",
                "input[placeholder*='codigo de acesso' i]",
                "button:has-text('Confirmar')",
            ]
            for sel in seletores_2fa:
                try:
                    if page.locator(sel).first.is_visible(timeout=150):
                        logger.warning(f"[2FA] Elemento de 2FA detectado via seletor '{sel}' no Playwright!")
                        return True
                except Exception:
                    pass

        # 3. Fallback de HTML raw
        try:
            conteudo = page.content().lower()
            if ("verificação em duas etapas" in conteudo or "verificacao em duas etapas" in conteudo) and \
               ("código de acesso" in conteudo or "codigo de acesso" in conteudo or "aplicativo gov.br" in conteudo):
                logger.warning("[2FA] Tela de 2FA detectada via HTML raw no Playwright!")
                return True
        except Exception:
            pass

    except Exception as e:
        logger.debug(f"Erro ao verificar tela de 2FA no Playwright: {e}")

    return False


def verificar_2fa_em_todas_paginas(context) -> tuple[bool, Optional[any]]:
    """Verifica se alguma das abas abertas no contexto exibe a tela de 2FA."""
    try:
        for p in context.pages:
            if not p.is_closed() and detectar_verificacao_duas_etapas_playwright(p):
                return True, p
    except Exception:
        pass
    return False, None


def clicar_botao_govbr_garantido(timeout_segundos: int = 25) -> bool:
    """
    Localiza e clica no botão 'Entrar com gov.br', garantindo que o clique
    realmente acionou a transição de tela.
    """
    fim = time.time() + timeout_segundos
    while time.time() < fim:
        pos = localizar_imagem(IMG_GOVBR, confidence=0.8)
        if pos:
            logger.info(f"[FASE 1 - SO] 'Entrar com gov.br' localizado em ({pos.x}, {pos.y}). Clicando...")
            clicar_robusto(pos.x, pos.y)
            time.sleep(2)

            # Verifica se o botão sumiu (indicando navegação iniciada)
            if not localizar_imagem(IMG_GOVBR, confidence=0.8):
                logger.info("[FASE 1 - SO] Botão 'Entrar com gov.br' clicado com sucesso!")
                return True
            else:
                logger.info("[FASE 1 - SO] Tentando segundo clique de ativação...")
                clicar_robusto(pos.x, pos.y)
                time.sleep(2)
                return True
        time.sleep(0.8)

    return False


def rolar_pagina_para_baixo_total():
    """
    Posiciona o mouse no centro da tela e rola a página completamente para baixo,
    garantindo que opções no rodapé (como 'Seu certificado digital') fiquem visíveis.
    """
    try:
        largura, altura = pyautogui.size()
        pyautogui.moveTo(largura // 2, altura // 2, duration=0.2)
        time.sleep(0.2)
        logger.info("[FASE 1 - SO] Realizando scroll para descer a página completamente...")
        for _ in range(5):
            pyautogui.scroll(-1000)
            time.sleep(0.1)
        time.sleep(0.5)
    except Exception as e:
        logger.warning(f"[FASE 1 - SO] Erro ao rolar a página para baixo: {e}")


def clicar_botao_certificado_garantido(timeout_segundos: int = 30) -> bool:
    """
    Aguarda e clica em 'Seu certificado digital' na tela do sso.acesso.gov.br.
    Necessariamente desce a tela toda antes de localizar e clicar na imagem do print.
    Se detectar captcha, levanta CaptchaDetectadoException imediatamente.
    """
    logger.info(f"[FASE 1 - SO] Preparando tela para 'Seu certificado digital' (timeout={timeout_segundos}s)...")

    # Aguarda renderização inicial da página do gov.br
    time.sleep(1.5)

    # Scroll obrigatório para descer a tela completamente ANTES de procurar o botão
    rolar_pagina_para_baixo_total()

    fim = time.time() + timeout_segundos
    tentativa = 0

    while time.time() < fim:
        tentativa += 1

        pos = localizar_imagem(IMG_CERT, confidence=0.75)
        if pos:
            logger.info(f"[FASE 1 - SO] 'Seu certificado digital' ENCONTRADO em ({pos.x}, {pos.y})! Clicando...")
            clicar_robusto(pos.x, pos.y)
            # Delay estendido para dar tempo ao Windows/Chrome de carregar a lista de certificados
            time.sleep(2.5)
            if detectar_presenca_captcha():
                raise CaptchaDetectadoException("hCaptcha detectado logo após clicar em 'Seu certificado digital'.")
            return True

        # Troca a verificação anterior: reforça o scroll para descer tudo caso a página ainda não tenha atingido o rodapé
        if tentativa % 5 == 0:
            logger.info(f"[FASE 1 - SO] Reforçando descida total da página (tentativa {tentativa})...")
            rolar_pagina_para_baixo_total()

        time.sleep(0.8)

    # Se falhar, salva captura para diagnóstico
    logger.warning(f"'Seu certificado digital' não foi localizado em {timeout_segundos}s.")
    try:
        pyautogui.screenshot(str(IMG_DEBUG_ERRO))
        logger.info(f"[DIAGNÓSTICO] Screenshot da tela salvo em: {IMG_DEBUG_ERRO.name}")
    except Exception:
        pass

    return False


def confirmar_popup_certificado(timeout_segundos: int = 25) -> bool:
    """
    Aguarda o popup nativo de seleção de certificado e clica em OK.
    Tenta por visão computacional (PyAutoGUI) e conta com fallbacks via UIA e ENTER.
    """
    logger.info("[FASE 1 - SO] Aguardando popup de seleção de certificado...")
    time.sleep(1.0)  # Delay para estabilização do popup na tela
    fim = time.time() + timeout_segundos
    tentativa = 0

    while time.time() < fim:
        tentativa += 1

        if detectar_presenca_captcha():
            raise CaptchaDetectadoException("hCaptcha detectado durante confirmação do popup de certificado.")

        # 1. Tenta por imagem (PyAutoGUI)
        pos_ok = localizar_imagem(IMG_OK, confidence=0.7)
        if pos_ok:
            logger.info(f"[FASE 1 - SO] Botão 'OK' do certificado encontrado em ({pos_ok.x}, {pos_ok.y})!")
            time.sleep(0.5)  # Delay antes do clique
            clicar_robusto(pos_ok.x, pos_ok.y)
            time.sleep(2.0)
            return True

        # 2. Fallback: via UIA nativo do Windows
        try:
            desktop = Desktop(backend="uia")
            for win in desktop.windows():
                titulo = win.window_text().lower()
                if any(ignorar in titulo for ignorar in [
                    "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                    "agendador de tarefas", "task scheduler"
                ]):
                    continue
                palavras_chave = [
                    "gov.br", "acesse sua conta", "selecionar um certificado",
                    "certificado", "segurança do windows", "windows security", "segurança"
                ]
                if any(k in titulo for k in palavras_chave):
                    btns_ok = win.descendants(control_type="Button", title="OK")
                    if btns_ok:
                        logger.info("[FASE 1 - SO] Botão 'OK' do certificado acionado via UIA!")
                        try:
                            focar_janela_segura(win)
                        except Exception:
                            pass
                        time.sleep(0.5)
                        btns_ok[0].click_input()
                        time.sleep(2.0)
                        return True
        except Exception:
            pass

        # A cada 4 tentativas (~3s), tenta focar a janela do certificado e enviar ENTER
        if tentativa % 4 == 0:
            logger.info("[FASE 1 - SO] Tentando confirmação intermediária com ENTER no popup...")
            try:
                desktop = Desktop(backend="uia")
                for win in desktop.windows():
                    titulo = win.window_text().lower()
                    if any(ignorar in titulo for ignorar in [
                        "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                        "agendador de tarefas", "task scheduler"
                    ]):
                        continue
                    if any(k in titulo for k in ["certificado", "segurança", "gov.br"]):
                        focar_janela_segura(win)
                        break
            except Exception:
                pass
            time.sleep(0.5)
            send_keys("{ENTER}")
            time.sleep(1.5)

        time.sleep(0.8)

    # 3. Fallback final: Enviar ENTER
    logger.info("[FASE 1 - SO] Enviando tecla ENTER para confirmar certificado...")
    time.sleep(0.8)
    send_keys("{ENTER}")
    time.sleep(2.0)
    return True


def selecionar_e_confirmar_certificado(termo_busca: str = "", timeout_segundos: int = 25) -> bool:
    """
    Localiza o popup nativo de certificado do Chrome, seleciona o certificado desejado
    pelo nome completo, CNPJ ou termo fornecido (suporta busca e rolagem para listas longas),
    salva screenshot de diagnóstico e confirma com OK/ENTER.
    """
    termo_str = str(termo_busca or "").strip()
    logger.info(f"[FASE 1 - SO] Aguardando popup de certificado para selecionar: '{termo_str or 'Padrão'}'...")
    time.sleep(1.0)
    fim = time.time() + timeout_segundos
    tentativa = 0

    # Extrai CNPJ (14 dígitos) do termo de busca se presente
    cnpj_busca = re.sub(r"\D", "", termo_str)
    if len(cnpj_busca) != 14:
        cnpj_busca = ""

    while time.time() < fim:
        tentativa += 1

        if detectar_presenca_captcha():
            raise CaptchaDetectadoException("hCaptcha detectado durante o aguardo do popup de certificado.")

        if detectar_presenca_2fa_uia():
            raise VerificacaoDuasEtapasException("Tela de Verificação em duas etapas (2FA) detectada.")

        try:
            desktop = Desktop(backend="uia")
            janela_cert = None
            for win in desktop.windows():
                titulo = win.window_text().lower()
                if any(ignorar in titulo for ignorar in [
                    "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                    "agendador de tarefas", "task scheduler"
                ]):
                    continue
                palavras_chave = [
                    "gov.br", "acesse sua conta", "selecionar um certificado",
                    "selecione um certificado", "certificado", "segurança do windows", "windows security", "segurança"
                ]
                if any(k in titulo for k in palavras_chave):
                    janela_cert = win
                    break

            if janela_cert:
                focar_janela_segura(janela_cert)
                logger.info(f"[CERTIFICADO] Janela de certificado localizada: '{janela_cert.window_text()}'")

                if not termo_str:
                    logger.info("[CERTIFICADO] Nenhum certificado específico informado na planilha; mantendo seleção padrão.")
                else:
                    item_alvo = None

                    def buscar_item_na_arvore():
                        for el in janela_cert.descendants():
                            try:
                                nome = el.window_text() or el.element_info.name or ""
                                nome_limpo = nome.strip()
                                if not nome_limpo:
                                    continue
                                # Critério A: CNPJ de 14 dígitos idêntico
                                if cnpj_busca and cnpj_busca in nome_limpo:
                                    return el
                                # Critério B: Nome completo ou termo contido no elemento
                                if termo_str.lower() in nome_limpo.lower() or nome_limpo.lower() in termo_str.lower():
                                    return el
                            except Exception:
                                pass
                        return None

                    item_alvo = buscar_item_na_arvore()

                    # Se não encontrou de imediato, pode ser uma lista longa com muitos certificados instalados
                    if not item_alvo:
                        logger.info(f"[CERTIFICADO] Certificado '{termo_str}' não visível de imediato. Percorrendo lista com rolagem (seta para baixo)...")
                        send_keys("{HOME}")
                        time.sleep(0.4)
                        for passo in range(50):  # Suporta até 50 certificados
                            item_alvo = buscar_item_na_arvore()
                            if item_alvo:
                                logger.info(f"[CERTIFICADO] Certificado localizado após rolagem no passo {passo + 1}!")
                                break
                            send_keys("{DOWN}")
                            time.sleep(0.15)

                    if item_alvo:
                        logger.info(f"[CERTIFICADO] Item correspondente encontrado via UIA: '{item_alvo.window_text()}'! Clicando...")
                        try:
                            item_alvo.click_input()
                            time.sleep(0.8)
                        except Exception as e_click:
                            logger.warning(f"[CERTIFICADO] Falha no click_input direto: {e_click}. Tentando via ENTER...")
                    else:
                        logger.error(
                            f"[CERTIFICADO] ERRO: Certificado '{termo_str}' NÃO foi localizado no Windows (verificados até 50 itens). "
                            "Cancelando seleção para evitar login com certificado incorreto."
                        )
                        try:
                            btns_cancelar = janela_cert.descendants(control_type="Button", title="Cancelar")
                            if btns_cancelar:
                                btns_cancelar[0].click_input()
                            else:
                                send_keys("{ESC}")
                        except Exception:
                            send_keys("{ESC}")
                        time.sleep(1.5)
                        return False

                # Confirma clicando no botão OK via UIA ou Imagem
                btns_ok = janela_cert.descendants(control_type="Button", title="OK")
                if btns_ok:
                    logger.info("[CERTIFICADO] Clicando no botão 'OK' via UIA...")
                    time.sleep(0.5)
                    btns_ok[0].click_input()
                    time.sleep(2.0)
                    return True

                pos_ok = localizar_imagem(IMG_OK, confidence=0.7)
                if pos_ok:
                    logger.info(f"[CERTIFICADO] Clicando no botão 'OK' via imagem ({pos_ok.x}, {pos_ok.y})...")
                    time.sleep(0.5)
                    clicar_robusto(pos_ok.x, pos_ok.y)
                    time.sleep(2.0)
                    return True

                logger.info("[CERTIFICADO] Confirmando via tecla ENTER...")
                time.sleep(0.5)
                send_keys("{ENTER}")
                time.sleep(2.0)
                return True

        except Exception as e:
            logger.warning(f"[CERTIFICADO] Tentativa {tentativa} falhou: {e}")

        time.sleep(0.8)

    # Fallback final se o loop acima esgotar
    if termo_str:
        logger.error(f"[CERTIFICADO] Timeout ao aguardar popup para seleção do certificado '{termo_str}'. Abortando.")
        return False

    logger.info("[FASE 1 - SO] Fallback: chamando confirmar_popup_certificado padrão...")
    return confirmar_popup_certificado(timeout_segundos=10)


def fechar_navegador(browser=None, context=None, page=None):
    """
    Fecha a janela do Chrome real e encerra todas as conexões
    de forma limpa e garantida, seja via Playwright, CDP, UIA ou processo.
    """
    global CHROME_PROC

    # 1. Encerramento gracioso via Playwright / CDP se ativos
    try:
        if context and page and not page.is_closed():
            cdp = context.new_cdp_session(page)
            cdp.send("Browser.close")
            time.sleep(0.5)
    except Exception:
        pass

    try:
        if browser:
            browser.close()
            time.sleep(0.3)
    except Exception:
        pass

    # 2. Encerramento forçado e definitivo de todos os processos da automação (árvore completa)
    matar_processos_chrome_automacao()

    # 3. Encerramento via fechamento de janela UIA se ainda existir resíduo
    try:
        desktop = Desktop(backend="uia")
        for win in desktop.windows():
            cls = win.class_name()
            txt = win.window_text()
            txt_lower = txt.lower()
            if any(ignorar in txt_lower for ignorar in [
                "antigravity", "visual studio code", "vscode", "microsoft edge", "msedge",
                "agendador de tarefas", "task scheduler"
            ]):
                continue
            if cls == "Chrome_WidgetWin_1" and ("google chrome" in txt_lower or txt_lower.endswith("chrome")):
                win.close()
    except Exception:
        pass

    # Aguarda liberação completa da porta CDP
    for _ in range(6):
        if not porta_em_uso(CDP_PORT):
            break
        time.sleep(0.5)

    time.sleep(0.5)


def limpar_ou_criar_pasta(caminho: Path):
    if not caminho.exists():
        caminho.mkdir(parents=True, exist_ok=True)
    else:
        for item in caminho.iterdir():
            try:
                if item.is_file() or item.is_symlink():
                    item.unlink()
                elif item.is_dir():
                    shutil.rmtree(item)
            except Exception as e:
                logger.warning(f"Não foi possível limpar {item}: {e}")


def remover_pastas_vazias_do_item(pasta_item: Path):
    """Replica a limpeza de subpastas vazias e da pasta principal do item."""
    for sub in ["CaixaPostal", "Notificacao", "ArquivosBaixados"]:
        sub_path = pasta_item / sub
        if sub_path.exists() and not any(sub_path.iterdir()):
            try:
                sub_path.rmdir()
            except OSError:
                pass

    if pasta_item.exists() and not any(pasta_item.iterdir()):
        try:
            pasta_item.rmdir()
        except OSError:
            pass


def obter_data_criacao(item: Path) -> datetime:
    """
    Retorna a data de criação do arquivo ou pasta.
    Tenta primeiro interpretar o timestamp no nome (ex: 23-10-2026_11-44-10 ou log_execucao_2026-10-23_11-44-10.txt),
    e como fallback utiliza a data de criação/modificação do sistema de arquivos.
    """
    nome = item.stem
    padroes = [
        r"(\d{2}-\d{2}-\d{4}_\d{2}-\d{2}-\d{2})",  # 23-10-2026_11-44-10
        r"(\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2})",  # 2026-10-23_11-44-10
    ]
    for padrao in padroes:
        m = re.search(padrao, nome)
        if m:
            texto_data = m.group(1)
            for fmt in ("%d-%m-%Y_%H-%M-%S", "%Y-%m-%d_%H-%M-%S"):
                try:
                    return datetime.strptime(texto_data, fmt)
                except ValueError:
                    pass

    stat = item.stat()
    timestamp = getattr(stat, "st_birthtime", None) or stat.st_ctime
    return datetime.fromtimestamp(timestamp)


def executar_rotatividade_antigos(pasta: Path, eh_pasta: bool, min_itens: int = 10, limite_dias: int = 60):
    """
    Sempre que houver pelo menos `min_itens` (10 itens) na pasta,
    exclui os itens mais antigos caso sua data de criação seja 60 dias ou mais anterior à data atual.
    """
    if not pasta.exists():
        return

    try:
        if eh_pasta:
            itens = [p for p in pasta.iterdir() if p.is_dir()]
        else:
            itens = [f for f in pasta.iterdir() if f.is_file() and f.suffix.lower() in (".txt", ".xlsx")]

        if len(itens) < min_itens:
            return

        itens_com_data = []
        for it in itens:
            try:
                dt = obter_data_criacao(it)
                itens_com_data.append((it, dt))
            except Exception:
                pass

        # Ordena do mais antigo para o mais recente
        itens_com_data.sort(key=lambda x: x[1])

        agora = datetime.now()
        removidos = 0

        for it, dt in itens_com_data:
            dias_idade = (agora - dt).days
            if dias_idade >= limite_dias:
                try:
                    if eh_pasta:
                        shutil.rmtree(it)
                        logger.info(f"[LIMPEZA HISTÓRICO] Pasta de capturas antiga removida: {it.name} ({dias_idade} dias de criação).")
                    else:
                        it.unlink()
                        logger.info(f"[LIMPEZA HISTÓRICO] Arquivo de log antigo removido: {it.name} ({dias_idade} dias de criação).")
                    removidos += 1
                except Exception as e_del:
                    logger.warning(f"[LIMPEZA HISTÓRICO] Não foi possível remover {it.name}: {e_del}")

        if removidos > 0:
            logger.info(f"[LIMPEZA HISTÓRICO] Total de {removidos} item(ns) com mais de {limite_dias} dias removido(s) de '{pasta.name}'.")
    except Exception as e:
        logger.warning(f"[LIMPEZA HISTÓRICO] Erro ao verificar rotatividade em '{pasta}': {e}")


def exibir_alerta_planilha_inexistente(caminho_criado: Path):
    mensagem = (
        "A Planilha de referência para consulta de CNPJs não existia.\n"
        f"A planilha modelo foi criada automaticamente no caminho:\n{caminho_criado}\n"
        "Por favor, preencha com os CNPJs na coluna 'CNPJ' e o certificado na coluna 'Certificado', e execute o robô novamente."
    )
    logger.warning(mensagem)


def efetuar_logout_det(page):
    """Efetua logout na sessão atual do DET clicando no botão Sair e confirmando."""
    if not page:
        return
    try:
        if page.is_closed():
            return
        time.sleep(DELAY_ENTRE_ACOES_S)
        logger.info("Efetuando logout no sistema DET...")

        btn_sair = page.locator("br-button:has-text('Sair'), button:has-text('Sair'), a:has-text('Sair')").first
        if not btn_sair.is_visible():
            btn_sair = page.get_by_text(re.compile(r"^\s*Sair\s*$", re.IGNORECASE)).first

        if btn_sair.is_visible():
            time.sleep(DELAY_ENTRE_ACOES_S)
            btn_sair.click()
            logger.info("Clique no botão 'Sair' realizado.")
            time.sleep(DELAY_ENTRE_ACOES_S)

            btn_confirmar = page.locator("button:has-text('Sim'), button:has-text('Confirmar'), br-button:has-text('Sim')").first
            try:
                if btn_confirmar.is_visible(timeout=2500):
                    btn_confirmar.click()
                    time.sleep(DELAY_ENTRE_ACOES_S)
            except Exception:
                pass

            time.sleep(1.5)
        else:
            logger.info("Botão 'Sair' não localizado no cabeçalho.")
    except Exception as e:
        logger.warning(f"Aviso durante tentativa de logout: {e}")


def processar_item_det(
    page,
    item_id: str,
    idx_item: int,
    total_itens: int,
    screenshots_dir: Path,
    hoje_date,
    perfil_ref: list,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    stop_event: Optional[threading.Event] = None,
) -> tuple[bool, str, str]:
    """
    Processa um único CNPJ dentro de uma sessão autenticada do DET:
    1. Troca de perfil para o procurador e digita o CNPJ.
    2. Valida se o CNPJ é aceito ou se há vínculo/procuração.
    3. Abre Caixa Postal, aplica filtro de não lidas e tira prints.
    4. Abre Notificações, confere datas, tira prints e baixa anexos.
    5. Retorna à tela de Serviços.
    """
    etapa = "Início"
    try:
        etapa = "Lendo o item da planilha"
        log_etapa(item_id, etapa, "Iniciando o processamento do item")

        etapa = "Aguardando a tela de serviços"
        page.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))").wait_for(state="visible", timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)

        etapa = "Trocando de perfil"
        btn_trocar_perfil = page.locator("button:has-text('Trocar Perfil')").first
        btn_trocar_perfil.wait_for(state="visible", timeout=20000)
        time.sleep(DELAY_ENTRE_ACOES_S)
        btn_trocar_perfil.click()
        time.sleep(2)

        # Clica no dropdown de seleção de perfil e navega via teclado como no PAD
        dropdown_perfil = page.locator("ng-select, .br-select, span.ng-arrow-wrapper").first
        dropdown_perfil.click()
        time.sleep(DELAY_ENTRE_ACOES_S)

        if perfil_ref[0] == 0:
            perfil_ref[0] += 1
            page.keyboard.press("ArrowDown")
            time.sleep(DELAY_ENTRE_ACOES_S)
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.press("Enter")
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.press("Tab")
        else:
            page.keyboard.press("ArrowDown")
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.press("ArrowDown")
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.press("Enter")
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.keyboard.press("Tab")

        time.sleep(DELAY_ENTRE_ACOES_S)
        page.keyboard.type(item_id, delay=100)
        time.sleep(DELAY_ENTRE_ACOES_S)

        etapa = "Selecionando o perfil e validando CNPJ"
        page.locator("button:has-text('Selecionar')").first.click()
        time.sleep(2)

        # 1. Verifica se o DET rejeitou por CNPJ inválido
        msg_invalido = page.get_by_text(re.compile(r"CNPJ\s+inv.lido", re.IGNORECASE))
        if msg_invalido.count() > 0 and msg_invalido.first.is_visible():
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.locator("button:has-text('Cancelar')").first.click()
            time.sleep(DELAY_ENTRE_ACOES_S)
            logger.warning(f"CNPJ recusado pelo DET, item ignorado: {item_id}")
            if progress_callback:
                progress_callback(idx_item, total_itens, f"Item {idx_item}/{total_itens} inválido no DET")
            return False, "CNPJ_RECUSADO_DET", "DET indicou CNPJ inválido"

        # 2. Verifica se o DET acusou falta de procuração ou vínculo
        msg_sem_procuracao = page.get_by_text(re.compile(r"(procura..o|v.nculo|n.o\s+encontrado|n.o\s+autorizado|perfil\s+n.o)", re.IGNORECASE))
        if msg_sem_procuracao.count() > 0 and msg_sem_procuracao.first.is_visible():
            time.sleep(DELAY_ENTRE_ACOES_S)
            page.locator("button:has-text('Cancelar')").first.click()
            time.sleep(DELAY_ENTRE_ACOES_S)
            logger.warning(f"Sem procuração ou vínculo no DET para o CNPJ: {item_id}")
            if progress_callback:
                progress_callback(idx_item, total_itens, f"Item {idx_item}/{total_itens} sem procuração")
            return False, "SEM_PROCURACAO", "Não foi encontrado procurador/vínculo no DET"

        # Criação de pastas do item
        etapa = "Criando a pasta do item"
        pasta_item = screenshots_dir / item_id
        pasta_caixa_postal = pasta_item / "CaixaPostal"
        pasta_notificacao = pasta_item / "Notificacao"
        pasta_arquivos_baixados = pasta_item / "ArquivosBaixados"

        pasta_caixa_postal.mkdir(parents=True, exist_ok=True)
        pasta_notificacao.mkdir(parents=True, exist_ok=True)
        pasta_arquivos_baixados.mkdir(parents=True, exist_ok=True)
        time.sleep(DELAY_ENTRE_ACOES_S)

        # --- CAIXA POSTAL ---
        etapa = "Abrindo a Caixa Postal"
        time.sleep(DELAY_ENTRE_ACOES_S)
        card_cp = page.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))")
        card_cp.wait_for(state="visible", timeout=30000)
        card_cp.click()
        time.sleep(2)
        formatted_dt = datetime.now().strftime("%d%m%Y_%H%M%S")

        # Aguarda a tela da Caixa Postal carregar
        try:
            page.locator("h1:has-text('Caixa Postal'), h2:has-text('Caixa Postal'), span:has-text('Caixa de Entrada')").first.wait_for(state="visible", timeout=15000)
        except Exception:
            pass
        time.sleep(DELAY_ENTRE_ACOES_S)

        # 1. Print ANTES de clicar em 'Exibir apenas não lidas' (visão geral / todas)
        etapa = "Tirando print da Caixa Postal antes do filtro"
        try:
            page.mouse.move(0, 0)
            time.sleep(0.3)
            page.evaluate("""() => {
                const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                tooltips.forEach(t => t.remove());
            }""")
        except Exception:
            pass

        caminho_print_antes = pasta_caixa_postal / f"{item_id}_{formatted_dt}_parte1_todas.jpg"
        page.screenshot(path=str(caminho_print_antes), type="jpeg")
        log_etapa(item_id, etapa, f"Print da Caixa Postal (antes do filtro / todas) realizado com sucesso: {caminho_print_antes.name}")
        time.sleep(DELAY_ENTRE_ACOES_S)

        # 2. Clica em 'Exibir apenas não lidas'
        etapa = "Filtrando por mensagens não lidas"
        loc_nao_lidas_filtro = page.locator(
            "span:has-text('Exibir apenas não lidas'), "
            "a:has-text('Exibir apenas não lidas'), "
            "button:has-text('Exibir apenas não lidas')"
        ).first

        clicou_filtro = False
        try:
            loc_nao_lidas_filtro.wait_for(state="visible", timeout=8000)
            time.sleep(0.3)
            loc_nao_lidas_filtro.click()
            clicou_filtro = True
            log_etapa(item_id, etapa, "Filtro 'Exibir apenas não lidas' acionado com sucesso.")
            time.sleep(2)
        except Exception:
            try:
                btn_fallback = page.get_by_text("Exibir apenas não lidas", exact=False).first
                if btn_fallback.is_visible():
                    btn_fallback.click()
                    clicou_filtro = True
                    log_etapa(item_id, etapa, "Filtro 'Exibir apenas não lidas' acionado via get_by_text.")
                    time.sleep(2)
            except Exception:
                pass

        if clicou_filtro:
            try:
                page.mouse.move(0, 0)
            except Exception:
                pass
        else:
            btn_todas = page.locator("span:has-text('Exibir todas'), a:has-text('Exibir todas'), button:has-text('Exibir todas')").first
            if btn_todas.is_visible():
                log_etapa(item_id, etapa, "Filtro de não lidas já está ativo ('Exibir todas' presente).")
            else:
                log_etapa(item_id, etapa, "Aviso: opção 'Exibir apenas não lidas' não foi localizada.")

        # 3. Depois de filtrar: verifica se há mensagens não lidas na tela
        etapa = "Tirando print da Caixa Postal após o filtro"
        time.sleep(0.5)

        loc_nao_lidas = page.locator(
            ".linha.nao-lida, "
            ".linha:has(i.fa-envelope.active), "
            ".linha:has(i.fa-envelope), "
            "i.fa-envelope.active, "
            "div[class*='nao-lida']"
        )
        loc_contador_aba = page.get_by_text(re.compile(r"Caixa de Entrada\s*\([1-9]\d*\)", re.I))
        linhas_tabela = page.locator(".tabela.mensagens .linha, .tabela .linha")
        msg_vazio = (
            page.locator(".empty-row")
            .or_(page.get_by_text(re.compile(r"nenhum registro|nenhuma mensagem|não há mensagens", re.I)))
        )
        btn_exibir_todas = page.locator("span:has-text('Exibir todas'), a:has-text('Exibir todas'), button:has-text('Exibir todas')")

        tem_msg_nao_lida = False
        try:
            if loc_nao_lidas.count() > 0 and loc_nao_lidas.first.is_visible():
                tem_msg_nao_lida = True
            elif loc_contador_aba.count() > 0 and loc_contador_aba.first.is_visible():
                tem_msg_nao_lida = True
            else:
                filtro_ativo = btn_exibir_todas.count() > 0 and btn_exibir_todas.first.is_visible()
                tabela_tem_linhas = linhas_tabela.count() > 0 and linhas_tabela.first.is_visible()
                tabela_vazia = msg_vazio.count() > 0 and msg_vazio.first.is_visible()

                if (filtro_ativo or clicou_filtro) and tabela_tem_linhas and not tabela_vazia:
                    tem_msg_nao_lida = True
                else:
                    try:
                        loc_nao_lidas.first.wait_for(state="visible", timeout=3000)
                        tem_msg_nao_lida = True
                    except PlaywrightTimeoutError:
                        tem_msg_nao_lida = False
        except Exception as e:
            logger.warning(f"Erro ao verificar mensagens não lidas: {e}")
            tem_msg_nao_lida = False

        # Tira SEMPRE o print após o filtro (mostrando apenas não lidas ou comprovando que não há nenhuma)
        time.sleep(DELAY_ENTRE_ACOES_S)
        try:
            page.mouse.move(0, 0)
            time.sleep(0.6)
            page.evaluate("""() => {
                const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                tooltips.forEach(t => t.remove());
            }""")
        except Exception:
            pass

        caminho_print_depois = pasta_caixa_postal / f"{item_id}_{formatted_dt}_parte2_nao_lidas.jpg"
        page.screenshot(path=str(caminho_print_depois), type="jpeg")
        if tem_msg_nao_lida:
            log_etapa(item_id, etapa, f"Print da Caixa Postal (após filtro / não lidas encontradas) realizado com sucesso: {caminho_print_depois.name}")
        else:
            log_etapa(item_id, etapa, f"Print da Caixa Postal (após filtro / sem não lidas) realizado com sucesso: {caminho_print_depois.name}")
        time.sleep(DELAY_ENTRE_ACOES_S)

        # Volta para a tela de Serviços
        time.sleep(DELAY_ENTRE_ACOES_S)
        btn_voltar = page.locator("a:has-text('Serviços'), button:has-text('Voltar'), i.fa-arrow-left, i.fa-home").first
        if btn_voltar.is_visible():
            btn_voltar.click()
        else:
            page.goto(URL_SERVICOS, timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)

        # --- NOTIFICAÇÃO ---
        etapa = "Abrindo a tela de Notificação"
        card_notificacao = page.locator(".cardListItem:has(span.title:has-text('NOTIFICAÇÃO'))")
        card_notificacao.wait_for(state="visible", timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)
        card_notificacao.click()
        time.sleep(DELAY_ENTRE_ACOES_S)

        etapa = "Tirando o print da Notificação"
        linhas_notif = page.locator("ngx-datatable datatable-row-wrapper")
        try:
            linhas_notif.first.wait_for(state="visible", timeout=10000)
            tem_notificacoes = True
        except PlaywrightTimeoutError:
            tem_notificacoes = False

        etapa = "Lendo a lista de notificações"
        if tem_notificacoes:
            qtd_linhas = linhas_notif.count()
            log_etapa(
                item_id,
                etapa,
                f"Lista de notificações lida: {qtd_linhas} linha(s) extraída(s). Data de hoje considerada: {hoje_date.strftime('%d/%m/%Y')}",
            )

            dados_linhas = []
            existe_notificacao_valida = False

            for idx in range(qtd_linhas):
                row_loc = linhas_notif.nth(idx)

                texto_data = ""
                todas_celulas = row_loc.locator("datatable-body-cell").all()
                for cell in todas_celulas:
                    c_txt = cell.inner_text().strip()
                    m = re.search(r"\b(\d{2}/\d{2}/\d{2,4})\b", c_txt)
                    if m:
                        texto_data = m.group(1)
                        break

                if not texto_data and len(todas_celulas) > 5:
                    try:
                        texto_data = todas_celulas[5].inner_text().strip()
                    except Exception:
                        pass

                data_notif = None
                valida = False
                for fmt in ("%d/%m/%Y", "%d/%m/%y"):
                    try:
                        data_notif = datetime.strptime(texto_data, fmt).date()
                        break
                    except ValueError:
                        pass

                if data_notif:
                    valida = (data_notif <= hoje_date) if MODO_TESTE else (data_notif >= hoje_date)

                if valida:
                    existe_notificacao_valida = True
                dados_linhas.append((idx, texto_data, valida))

            if not existe_notificacao_valida:
                try:
                    page.mouse.move(0, 0)
                    time.sleep(0.3)
                    page.evaluate("""() => {
                        const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                        tooltips.forEach(t => t.remove());
                    }""")
                except Exception:
                    pass

                caminho_print_notif = pasta_notificacao / f"{item_id}_{formatted_dt}.jpg"
                page.screenshot(path=str(caminho_print_notif), type="jpeg")
                log_etapa(
                    item_id,
                    etapa,
                    f"Print da tela de Notificação (nenhuma válida por data) realizado com sucesso: {caminho_print_notif.name}",
                )
                time.sleep(DELAY_ENTRE_ACOES_S)
            else:
                # --- CAPTURA DE PRINTS DE TELA INTEIRA (SEM DUPLICAÇÃO) ---
                etapa = "Tirando prints da tela de Notificações"
                indices_validos = [idx for idx, _, val in dados_linhas if val]
                linhas_ja_capturadas = set()

                # 1. Primeiro print: Tela inicial no topo (mostrando dados do Empregador, Usuário e filtros)
                try:
                    page.evaluate("window.scrollTo(0, 0)")
                except Exception:
                    pass
                time.sleep(DELAY_ENTRE_ACOES_S)

                # Remove tooltips / popovers antes do print
                try:
                    page.mouse.move(0, 0)
                    time.sleep(0.3)
                    page.evaluate("""() => {
                        const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                        tooltips.forEach(t => t.remove());
                    }""")
                except Exception:
                    pass

                caminho_print_1 = pasta_notificacao / f"{item_id}_{formatted_dt}_parte1.jpg"
                page.screenshot(path=str(caminho_print_1), type="jpeg")
                log_etapa(
                    item_id,
                    etapa,
                    f"Print #1 da tela inteira (topo com identificação do Empregador) realizado com sucesso ({caminho_print_1.name}).",
                )
                time.sleep(DELAY_ENTRE_ACOES_S)

                # 2. Segundo print: Rola até a primeira notificação (ou cabeçalho) para posicioná-la no topo
                try:
                    topo_notif = page.locator("h1:has-text('Notificações'), h2:has-text('Notificações'), h3:has-text('Notificações'), ngx-datatable").first
                    if topo_notif.count() > 0 and topo_notif.is_visible():
                        topo_notif.evaluate("el => el.scrollIntoView({ block: 'start' })")
                    elif linhas_notif.count() > 0:
                        linhas_notif.first.evaluate("el => el.scrollIntoView({ block: 'start' })")
                except Exception:
                    try:
                        linhas_notif.first.scroll_into_view_if_needed()
                    except Exception:
                        pass

                time.sleep(DELAY_ENTRE_ACOES_S)

                try:
                    page.mouse.move(0, 0)
                    time.sleep(0.3)
                    page.evaluate("""() => {
                        const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                        tooltips.forEach(t => t.remove());
                    }""")
                except Exception:
                    pass

                # Identifica quais linhas ficaram visíveis neste segundo print a partir do topo das notificações
                try:
                    linhas_visiveis = page.evaluate("""() => {
                        const rows = Array.from(document.querySelectorAll('ngx-datatable datatable-row-wrapper'));
                        const vh = window.innerHeight;
                        const vis = [];
                        rows.forEach((r, i) => {
                            const rect = r.getBoundingClientRect();
                            const hVis = Math.min(rect.bottom, vh) - Math.max(rect.top, 0);
                            if (hVis > 50 && rect.top < vh && rect.bottom > 0) {
                                vis.push(i);
                            }
                        });
                        return vis;
                    }""")
                    linhas_ja_capturadas.update(linhas_visiveis)
                except Exception:
                    linhas_ja_capturadas.add(0)

                # Tira o print com a tabela rolada para o topo
                caminho_print_2 = pasta_notificacao / f"{item_id}_{formatted_dt}_parte2.jpg"
                page.screenshot(path=str(caminho_print_2), type="jpeg")
                log_etapa(
                    item_id,
                    etapa,
                    f"Print #2 da tela inteira (rolado até as notificações) realizado com sucesso ({caminho_print_2.name}), cobrindo linha(s): {sorted(list(linhas_ja_capturadas))}.",
                )
                time.sleep(DELAY_ENTRE_ACOES_S)

                # 3. Se houver notificações válidas mais abaixo que ainda não apareceram no print rolado, vai descendo
                num_print = 3
                for idx in indices_validos:
                    if idx in linhas_ja_capturadas:
                        continue

                    row_loc = linhas_notif.nth(idx)
                    try:
                        row_loc.evaluate("el => el.scrollIntoView({ block: 'start' })")
                    except Exception:
                        row_loc.scroll_into_view_if_needed()
                    time.sleep(DELAY_ENTRE_ACOES_S)

                    try:
                        page.mouse.move(0, 0)
                        time.sleep(0.3)
                        page.evaluate("""() => {
                            const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                            tooltips.forEach(t => t.remove());
                        }""")
                    except Exception:
                        pass

                    try:
                        linhas_visiveis = page.evaluate("""() => {
                            const rows = Array.from(document.querySelectorAll('ngx-datatable datatable-row-wrapper'));
                            const vh = window.innerHeight;
                            const vis = [];
                            rows.forEach((r, i) => {
                                const rect = r.getBoundingClientRect();
                                const hVis = Math.min(rect.bottom, vh) - Math.max(rect.top, 0);
                                if (hVis > 50 && rect.top < vh && rect.bottom > 0) {
                                    vis.push(i);
                                }
                            });
                            return vis;
                        }""")
                        linhas_ja_capturadas.update(linhas_visiveis)
                    except Exception:
                        linhas_ja_capturadas.add(idx)

                    caminho_print_n = pasta_notificacao / f"{item_id}_{formatted_dt}_parte{num_print}.jpg"
                    page.screenshot(path=str(caminho_print_n), type="jpeg")
                    log_etapa(
                        item_id,
                        etapa,
                        f"Print #{num_print} da tela inteira realizado descendo até a linha {idx} ({caminho_print_n.name}), cobrindo linha(s): {sorted(list(linhas_ja_capturadas))}.",
                    )
                    num_print += 1
                    time.sleep(DELAY_ENTRE_ACOES_S)

            for idx, texto_data, linha_valida in dados_linhas:
                etapa = f"Lendo a data da linha {idx}"
                if linha_valida:
                    row_loc = linhas_notif.nth(idx)

                    etapa = f"Rolando a página até a linha {idx} para download"
                    row_loc.scroll_into_view_if_needed()
                    time.sleep(DELAY_ENTRE_ACOES_S)

                    etapa = f"Localizando o botão de download da linha {idx}"
                    candidatos_download = row_loc.locator(
                        "button:has(i.fa-download), a:has(i.fa-download), "
                        "button:has(i.fa-print), a:has(i.fa-print), "
                        "button[title*='Download'], button[title*='Baixar'], button[title*='Imprimir']"
                    )
                    btn_download = None
                    if candidatos_download.count() > 0:
                        btn_download = candidatos_download.first
                    else:
                        celula_acoes = row_loc.locator("datatable-body-cell").last
                        botoes_acoes = celula_acoes.locator("button, a")
                        if botoes_acoes.count() > 1:
                            btn_download = botoes_acoes.nth(1)
                        elif botoes_acoes.count() == 1:
                            btn_download = botoes_acoes.first

                    if btn_download and btn_download.count() > 0 and btn_download.is_visible():
                        log_etapa(item_id, etapa, f"Linha {idx} (data {texto_data}): clicando no botão de download")
                        try:
                            etapa = f"Aguardando o download da linha {idx}"
                            time.sleep(DELAY_ENTRE_ACOES_S)
                            with page.expect_download(timeout=60000) as download_info:
                                btn_download.click()
                                log_etapa(item_id, etapa, f"Linha {idx}: clique no botão de download realizado")

                            download = download_info.value
                            sug_name = download.suggested_filename
                            destino_final = pasta_arquivos_baixados / sug_name
                            if destino_final.exists():
                                stem = destino_final.stem
                                ext = destino_final.suffix
                                destino_final = pasta_arquivos_baixados / f"{stem}_linha_{idx}{ext}"

                            download.save_as(str(destino_final))
                            time.sleep(DELAY_ENTRE_ACOES_S)
                            log_etapa(
                                item_id,
                                etapa,
                                f"Linha {idx}: arquivo baixado e movido para ArquivosBaixados ({destino_final.name})",
                            )
                        except PlaywrightTimeoutError:
                            log_etapa(
                                item_id,
                                etapa,
                                f"AVISO: download não detectado em 60s (linha {idx})",
                            )
                    else:
                        log_etapa(
                            item_id,
                            etapa,
                            f"Linha {idx} (data {texto_data}): notificação válida sem botão de download (apenas leitura).",
                        )
                else:
                    log_etapa(
                        item_id,
                        etapa,
                        f"Linha {idx} ignorada: data {texto_data} fora do critério (ModoTeste={MODO_TESTE}, hoje={hoje_date.strftime('%d/%m/%Y')})",
                    )
        else:
            # Tabela vazia: tira o print de comprovação da tela de notificações vazia
            try:
                page.mouse.move(0, 0)
                time.sleep(0.3)
                page.evaluate("""() => {
                    const tooltips = document.querySelectorAll('.tooltip, [role="tooltip"], br-tooltip, .br-tooltip, div[class*="tooltip"], div[class*="popover"]');
                    tooltips.forEach(t => t.remove());
                }""")
            except Exception:
                pass

            caminho_print_notif = pasta_notificacao / f"{item_id}_{formatted_dt}.jpg"
            page.screenshot(path=str(caminho_print_notif), type="jpeg")
            log_etapa(item_id, etapa, f"Print da tela de Notificação (vazia / sem notificações) realizado com sucesso: {caminho_print_notif.name}")
            time.sleep(DELAY_ENTRE_ACOES_S)

        # Retorna à tela de serviços
        time.sleep(DELAY_ENTRE_ACOES_S)
        page.goto(URL_SERVICOS, timeout=60000)
        time.sleep(DELAY_ENTRE_ACOES_S)

        # Limpeza de diretórios vazios e contabilização
        etapa = "Limpando subpastas vazias"
        remover_pastas_vazias_do_item(pasta_item)

        arquivos_guardados = list(pasta_item.rglob("*")) if pasta_item.exists() else []
        qtd_arquivos = sum(1 for f in arquivos_guardados if f.is_file())

        logger.info(
            f"SUCESSO: O item '{item_id}' foi verificado. Arquivos guardados (prints + downloads): {qtd_arquivos}"
        )
        return True, "SUCESSO", f"{qtd_arquivos} arquivo(s) guardado(s)"

    except Exception as e:
        logger.error(f"Erro no item '{item_id}' na etapa '{etapa}': {e}")
        try:
            page.goto(URL_SERVICOS, timeout=60000)
        except Exception:
            pass
        return False, "FALHA_SISTEMA", f"Erro na etapa '{etapa}': {e}"


def processar_item_det_com_retry(
    page,
    item_id: str,
    idx_item: int,
    total_itens: int,
    screenshots_dir: Path,
    hoje_date,
    perfil_ref: list,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    stop_event: Optional[threading.Event] = None,
) -> tuple[bool, str, str]:
    """
    Executa processar_item_det e, em caso de FALHA_SISTEMA temporária,
    efetua 1 retry imediato retornando para a tela de Serviços.
    """
    sucesso, status_code, detalhe = processar_item_det(
        page=page,
        item_id=item_id,
        idx_item=idx_item,
        total_itens=total_itens,
        screenshots_dir=screenshots_dir,
        hoje_date=hoje_date,
        perfil_ref=perfil_ref,
        progress_callback=progress_callback,
        stop_event=stop_event,
    )

    if not sucesso and status_code == "FALHA_SISTEMA":
        if stop_event and stop_event.is_set():
            return sucesso, status_code, detalhe
        logger.warning(f"Tentando reprocessamento imediato (1 retry) do CNPJ '{item_id}'...")
        try:
            page.goto(URL_SERVICOS, timeout=45000)
            time.sleep(2)
            sucesso, status_code, detalhe = processar_item_det(
                page=page,
                item_id=item_id,
                idx_item=idx_item,
                total_itens=total_itens,
                screenshots_dir=screenshots_dir,
                hoje_date=hoje_date,
                perfil_ref=perfil_ref,
                progress_callback=progress_callback,
                stop_event=stop_event,
            )
            if sucesso:
                detalhe = f"{detalhe} (recuperado no retry imediato)"
        except Exception as e_ret:
            detalhe = f"Falha persistente após retry imediato: {e_ret}"

    return sucesso, status_code, detalhe


class ItemExecucao:
    """Representa o ciclo de vida e histórico de execução de um CNPJ nas duas tentativas."""
    def __init__(self, cnpj: str, certificado: str):
        self.cnpj = str(cnpj).strip()
        self.certificado = str(certificado).strip()
        self.status_t1 = "NAO_EXECUTADO"
        self.detalhe_t1 = ""
        self.status_t2 = "-"
        self.detalhe_t2 = "-"
        self.status_final = "NAO_EXECUTADO"
        self.data_hora = ""


def salvar_relatorio_execucao(base_dir: Path, execucao_dt_str: str, itens: list[ItemExecucao]) -> Path:
    """
    Cria a pasta Relatorios se não existir e gera o arquivo Excel com o histórico
    detalhado de todas as tentativas e status consolidado de cada CNPJ.
    """
    pasta_relatorios = base_dir / "Relatorios"
    pasta_relatorios.mkdir(parents=True, exist_ok=True)
    caminho_relatorio = pasta_relatorios / f"Relatorio_Execucao_{execucao_dt_str}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório de Execução"

    headers = [
        "CNPJ",
        "Certificado",
        "Status 1ª Tentativa",
        "Detalhe 1ª Tentativa",
        "Status 2ª Tentativa",
        "Detalhe 2ª Tentativa",
        "Status Final",
        "Data/Hora",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1351B4", end_color="1351B4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    fill_sucesso = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_2fa = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_falha = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for it in itens:
        status_final_texto = it.status_final
        if it.status_final == "EXIGE_2FA":
            status_final_texto = "Necessita confirmação de identidade"
        elif it.status_final == "SUCESSO":
            status_final_texto = "SUCESSO"

        ws.append([
            it.cnpj,
            it.certificado,
            it.status_t1,
            it.detalhe_t1,
            it.status_t2,
            it.detalhe_t2,
            status_final_texto,
            it.data_hora or datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        ])
        row_idx = ws.max_row
        status_cell = ws.cell(row=row_idx, column=7)
        if it.status_final == "SUCESSO":
            status_cell.fill = fill_sucesso
        elif it.status_final == "EXIGE_2FA":
            status_cell.fill = fill_2fa
        elif it.status_final in ("FALHA_SISTEMA", "NAO_EXECUTADO"):
            status_cell.fill = fill_falha

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(str(caminho_relatorio))
    wb.close()
    logger.info(f"[RELATÓRIO] Planilha de relatório geral gerada com sucesso em: {caminho_relatorio.name}")
    return caminho_relatorio





def salvar_planilha_pendentes(base_dir: Path, itens: list[ItemExecucao]) -> Optional[Path]:
    """
    Gera a planilha CNPJs_Pendentes.xlsx contendo os CNPJs com falha temporária,
    não executados ou bloqueados por 2FA. Caso não haja pendentes, remove qualquer versão residual anterior.
    """
    caminho_pendentes = base_dir / "CNPJs_Pendentes.xlsx"
    pendentes = [it for it in itens if it.status_final in ("FALHA_SISTEMA", "NAO_EXECUTADO", "EXIGE_2FA")]

    if pendentes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CNPJs_consultados"
        ws["A1"] = "CNPJ"
        ws["B1"] = "Certificado"

        cor_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        fonte_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        alinhamento_centro = Alignment(horizontal="center", vertical="center")
        borda_fina = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        ws.row_dimensions[1].height = 24
        for col_idx in [1, 2]:
            c = ws.cell(row=1, column=col_idx)
            c.fill = cor_header
            c.font = fonte_header
            c.alignment = alinhamento_centro

        for p in pendentes:
            ws.append([p.cnpj, p.certificado])

        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 20
            c_cnpj = ws.cell(row=row_idx, column=1)
            c_cert = ws.cell(row=row_idx, column=2)
            c_cnpj.font = Font(name="Calibri", size=11)
            c_cnpj.alignment = alinhamento_centro
            c_cnpj.border = borda_fina
            c_cnpj.number_format = "@"
            c_cert.font = Font(name="Calibri", size=11)
            c_cert.alignment = Alignment(horizontal="left", vertical="center")
            c_cert.border = borda_fina

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 65

        wb.save(str(caminho_pendentes))
        wb.close()
        logger.info(f"[PENDENTES] Planilha 'CNPJs_Pendentes.xlsx' gerada com {len(pendentes)} CNPJ(s) pendente(s).")
        return caminho_pendentes
    else:
        if caminho_pendentes.exists():
            try:
                caminho_pendentes.unlink()
                logger.info("[PENDENTES] Todos os CNPJs foram processados com sucesso. Arquivo 'CNPJs_Pendentes.xlsx' anterior removido.")
            except Exception as e:
                logger.warning(f"[PENDENTES] Não foi possível remover planilha anterior de pendentes: {e}")
        return None


def registrar_pulo_2fa(
    lista_cnpjs: list[str],
    mapa_itens: dict[str, ItemExecucao],
    nome_cert_exibicao: str,
    rodada: int,
    num_grupo: int,
    total_grupos: int,
    total_itens: int,
    item_global_counter: list[int],
    pasta_cert: Path,
    base_dir: Path,
    execucao_dt_str: str,
    pagina_2fa=None,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
):
    """
    Registra o pulo imediato de um certificado que exibiu a tela de Verificação em duas etapas (2FA).
    Salva screenshot comprobatório na pasta do certificado, preenche os itens como
    'Necessita confirmação de identidade' e grava imediatamente nos relatórios e pendentes.
    """
    logger.warning("\n" + "=" * 80)
    logger.warning(
        f"[VERIFICAÇÃO EM DUAS ETAPAS] O certificado '{nome_cert_exibicao}' exibiu tela de Verificação em duas etapas no Gov.br.\n"
        f"Fechando o navegador e pulando todos os {len(lista_cnpjs)} CNPJ(s) associados a este certificado."
    )
    logger.warning("=" * 80 + "\n")

    pasta_cert.mkdir(parents=True, exist_ok=True)
    caminho_print_2fa = pasta_cert / "01_necessita_confirmacao_identidade_2fa.png"
    try:
        salvo = False
        if pagina_2fa and not pagina_2fa.is_closed():
            try:
                pagina_2fa.bring_to_front()
                pagina_2fa.screenshot(path=str(caminho_print_2fa))
                salvo = True
            except Exception:
                pass
        if not salvo:
            try:
                pyautogui.screenshot(str(caminho_print_2fa))
                salvo = True
            except Exception:
                pass
        if not salvo:
            img = Image.new("RGB", (1280, 720), color=(255, 242, 204))
            img.save(str(caminho_print_2fa))
            salvo = True
        logger.info(f"[COMPROVANTE] Screenshot da tela de 2FA salvo na pasta do certificado: {caminho_print_2fa.name}")
    except Exception as e_shot:
        logger.warning(f"Não foi possível salvar screenshot de 2FA: {e_shot}")

    for doc in lista_cnpjs:
        it = mapa_itens[doc]
        status_code = "EXIGE_2FA"
        status_exibicao = "Necessita confirmação de identidade"
        msg = "Necessita confirmação de identidade (2FA Gov.br)"
        if rodada == 1:
            it.status_t1 = status_exibicao
            it.detalhe_t1 = msg
        else:
            it.status_t2 = status_exibicao
            it.detalhe_t2 = msg
        it.status_final = status_code
        it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    item_global_counter[0] += len(lista_cnpjs)

    # Grava imediatamente no Relatório de Execução e na Planilha de Pendentes
    try:
        salvar_relatorio_execucao(base_dir, execucao_dt_str, list(mapa_itens.values()))
        salvar_planilha_pendentes(base_dir, list(mapa_itens.values()))
    except Exception as e_rel:
        logger.debug(f"Aviso ao salvar relatórios após 2FA: {e_rel}")

    if progress_callback:
        progress_callback(
            item_global_counter[0],
            total_itens,
            f"Grupo {num_grupo}/{total_grupos} pulado: 2FA ({nome_cert_exibicao[:20]})",
        )


# ==============================================================================
# FLUXO PRINCIPAL DA AUTOMAÇÃO HÍBRIDA
# ==============================================================================
def processar_sessao_certificado(
    cert_alvo: str,
    lista_cnpjs: list[str],
    mapa_itens: dict[str, ItemExecucao],
    rodada: int,
    num_grupo: int,
    total_grupos: int,
    total_itens: int,
    item_global_counter: list[int],
    screenshots_dir: Path,
    hoje_date,
    base_dir: Path,
    execucao_dt_str: str,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    stop_event: Optional[threading.Event] = None,
) -> bool:
    """
    Executa a sessão completa para um grupo de certificado (Fases 1 a 4):
    - Login no Gov.br com retry contra Captcha e Fallback com espera de 60s em caso de portal fora do ar.
    - Conexão Playwright via CDP.
    - Processamento de cada CNPJ com retry imediato em caso de erro temporário.
    - Atualização do histórico em mapa_itens (rodada 1 ou rodada 2).
    - Logout e fechamento limpo do navegador.
    """
    nome_cert_exibicao = cert_alvo if cert_alvo else "Padrão"
    pasta_cert = screenshots_dir / sanitizar_nome_pasta(cert_alvo)
    pasta_cert.mkdir(parents=True, exist_ok=True)
    logger.info(
        f"\n{'='*60}\n>>> INICIANDO GRUPO {num_grupo}/{total_grupos} (Rodada {rodada}): "
        f"Certificado '{nome_cert_exibicao}' ({len(lista_cnpjs)} CNPJs)\n{'='*60}"
    )

    playwright_mgr = None
    browser = None
    context = None
    page = None

    try:
        # ==================================================================
        # FASE 1: Login com Retry contra Captcha e Fallback de 60s contra Gov.br Fora do Ar
        # ==================================================================
        MAX_TENTATIVAS_CAPTCHA = 5
        MAX_TENTATIVAS_PORTAL = 2
        sucesso_fase1 = False

        for tentativa_portal in range(1, MAX_TENTATIVAS_PORTAL + 1):
            if stop_event and stop_event.is_set():
                break

            falha_botoes_portal = False

            for tentativa_captcha in range(1, MAX_TENTATIVAS_CAPTCHA + 1):
                if stop_event and stop_event.is_set():
                    break

                logger.info(
                    f"\n>>> Fase 1 (Portal {tentativa_portal}/{MAX_TENTATIVAS_PORTAL} | "
                    f"Tentativa {tentativa_captcha}/{MAX_TENTATIVAS_CAPTCHA}): "
                    f"Abertura do Chrome e Login gov.br para '{nome_cert_exibicao}'..."
                )
                fechar_navegador(browser, context, page)
                liberar_porta_se_ocupada()
                iniciar_chrome_hibrido()

                try:
                    chrome_win = obter_janela_chrome(timeout_segundos=30)
                    focar_janela_segura(chrome_win)
                    time.sleep(2)

                    if stop_event and stop_event.is_set():
                        fechar_navegador(browser, context, page)
                        break

                    if not clicar_botao_govbr_garantido(timeout_segundos=25):
                        logger.warning("Não foi possível localizar o botão 'Entrar com gov.br'.")
                        falha_botoes_portal = True
                        fechar_navegador(browser, context, page)
                        break

                    if stop_event and stop_event.is_set():
                        fechar_navegador(browser, context, page)
                        break

                    if not clicar_botao_certificado_garantido(timeout_segundos=30):
                        logger.warning("Não foi possível localizar 'Seu certificado digital'.")
                        falha_botoes_portal = True
                        fechar_navegador(browser, context, page)
                        break

                    if stop_event and stop_event.is_set():
                        fechar_navegador(browser, context, page)
                        break

                    # Seleciona o certificado específico do grupo no popup nativo
                    sucesso_cert = selecionar_e_confirmar_certificado(termo_busca=cert_alvo, timeout_segundos=25)
                    if not sucesso_cert:
                        logger.error(
                            f"\n[ATENÇÃO] O certificado '{nome_cert_exibicao}' NÃO foi localizado ou não está instalado no Windows. "
                            f"Todos os {len(lista_cnpjs)} CNPJ(s) associados a este grupo foram ignorados."
                        )
                        for doc in lista_cnpjs:
                            it = mapa_itens[doc]
                            msg = f"Certificado '{nome_cert_exibicao}' não localizado no Windows"
                            if rodada == 1:
                                it.status_t1 = "NAO_EXECUTADO"
                                it.detalhe_t1 = msg
                            else:
                                it.status_t2 = "NAO_EXECUTADO"
                                it.detalhe_t2 = msg
                            it.status_final = "NAO_EXECUTADO"
                            it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

                        item_global_counter[0] += len(lista_cnpjs)
                        if progress_callback:
                            progress_callback(
                                item_global_counter[0],
                                total_itens,
                                f"Grupo {num_grupo}/{total_grupos} ignorado (Certificado não encontrado)",
                            )
                        fechar_navegador(browser, context, page)
                        return False

                    sucesso_fase1 = True
                    break  # Login com sucesso!

                except VerificacaoDuasEtapasException:
                    registrar_pulo_2fa(
                        lista_cnpjs=lista_cnpjs,
                        mapa_itens=mapa_itens,
                        nome_cert_exibicao=nome_cert_exibicao,
                        rodada=rodada,
                        num_grupo=num_grupo,
                        total_grupos=total_grupos,
                        total_itens=total_itens,
                        item_global_counter=item_global_counter,
                        pasta_cert=pasta_cert,
                        base_dir=base_dir,
                        execucao_dt_str=execucao_dt_str,
                        pagina_2fa=page,
                        progress_callback=progress_callback,
                    )
                    fechar_navegador(browser, context, page)
                    return False

                except CaptchaDetectadoException as e_cap:
                    logger.warning(
                        f"\n[CAPTCHA DETECTADO] {e_cap}\n"
                        f"Fechando o navegador e tentando novamente (tentativa {tentativa_captcha}/{MAX_TENTATIVAS_CAPTCHA})..."
                    )
                    fechar_navegador(browser, context, page)
                    time.sleep(3)
                    if tentativa_captcha == MAX_TENTATIVAS_CAPTCHA:
                        logger.error(f"[CAPTCHA] Limite de {MAX_TENTATIVAS_CAPTCHA} tentativas atingido para '{nome_cert_exibicao}'.")
                        falha_botoes_portal = True
                        break

            if sucesso_fase1:
                break

            # Se houve falha de portal (site fora do ar ou botões não surgiram):
            if falha_botoes_portal and tentativa_portal < MAX_TENTATIVAS_PORTAL:
                logger.warning("\n" + "=" * 80)
                logger.warning("[ATENÇÃO - NÃO MEXA NO COMPUTADOR]")
                logger.warning("O portal gov.br apresentou instabilidade ou lentidão na resposta.")
                logger.warning(f"O navegador será reiniciado e aguardará 60 segundos antes da tentativa {tentativa_portal + 1}/{MAX_TENTATIVAS_PORTAL}.")
                logger.warning("=" * 80 + "\n")
                fechar_navegador(browser, context, page)

                for seg_restantes in range(60, 0, -10):
                    if stop_event and stop_event.is_set():
                        break
                    logger.info(f"[AGUARDE] Tentando reconexão ao gov.br em {seg_restantes} segundos... (Não mexa no computador)")
                    if progress_callback:
                        progress_callback(
                            item_global_counter[0],
                            total_itens,
                            f"Gov.br instável. Aguardando {seg_restantes}s para tentar novamente...",
                        )
                    time.sleep(10)
            elif falha_botoes_portal and tentativa_portal >= MAX_TENTATIVAS_PORTAL:
                # Esgotou as 2 tentativas do portal
                logger.error(f"[PORTAL GOV.BR INDISPONÍVEL] O portal não respondeu após 2 tentativas com intervalo de 60s.")
                caminho_print_erro_gov = screenshots_dir / f"erro_govbr_indisponivel_{nome_cert_exibicao[:20]}.jpg"
                try:
                    pyautogui.screenshot(str(caminho_print_erro_gov))
                    logger.info(f"[COMPROVANTE] Screenshot do erro no gov.br salvo em: {caminho_print_erro_gov.name}")
                except Exception:
                    pass

                for doc in lista_cnpjs:
                    it = mapa_itens[doc]
                    msg = "Portal Gov.br indisponível após 2 tentativas (espera de 60s)"
                    if rodada == 1:
                        it.status_t1 = "NAO_EXECUTADO"
                        it.detalhe_t1 = msg
                    else:
                        it.status_t2 = "NAO_EXECUTADO"
                        it.detalhe_t2 = msg
                    it.status_final = "NAO_EXECUTADO"
                    it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

                item_global_counter[0] += len(lista_cnpjs)
                if progress_callback:
                    progress_callback(
                        item_global_counter[0],
                        total_itens,
                        f"Grupo {num_grupo}/{total_grupos} falhou: Portal Gov.br indisponível",
                    )
                fechar_navegador(browser, context, page)
                return False

        if not sucesso_fase1:
            fechar_navegador(browser, context, page)
            return False

        logger.info(f"Fase 1 concluída: Certificado '{nome_cert_exibicao}' autenticado. Conectando Playwright...")

        # ==================================================================
        # FASE 2: Conexão Playwright via CDP
        # ==================================================================
        playwright_mgr = sync_playwright().start()

        for tentativa_cdp in range(1, 6):
            if stop_event and stop_event.is_set():
                break
            try:
                browser = playwright_mgr.chromium.connect_over_cdp(
                    f"http://127.0.0.1:{CDP_PORT}",
                    slow_mo=DELAY_SLOW_MO_MS
                )
                break
            except Exception:
                time.sleep(2)

        if not browser:
            logger.error("Erro fatal: Não foi possível conectar o Playwright via CDP após 5 tentativas.")
            if playwright_mgr:
                playwright_mgr.stop()
            for doc in lista_cnpjs:
                it = mapa_itens[doc]
                msg = "Falha de conexão CDP do Playwright"
                if rodada == 1:
                    it.status_t1 = "FALHA_SISTEMA"
                    it.detalhe_t1 = msg
                else:
                    it.status_t2 = "FALHA_SISTEMA"
                    it.detalhe_t2 = msg
                it.status_final = "FALHA_SISTEMA"
                it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
            return False

        context = browser.contexts[0]
        page = context.pages[0] if context.pages else context.new_page()

        inicio_espera = time.time()
        timeout_login = 45
        sessao_estabelecida = False
        tela_2fa_detectada = False
        pagina_2fa = None

        logger.info(f"Aguardando autenticação e carregamento do portal DET para '{nome_cert_exibicao}'...")

        while time.time() - inicio_espera < timeout_login:
            if stop_event and stop_event.is_set():
                break

            # 1. Verifica se a tela de 2FA surgiu em alguma das páginas do navegador
            tem_2fa, p_2fa = verificar_2fa_em_todas_paginas(context)
            if not tem_2fa and detectar_presenca_2fa_uia():
                tem_2fa = True
                p_2fa = page

            if tem_2fa:
                tela_2fa_detectada = True
                pagina_2fa = p_2fa
                break

            # 2. Verifica se houve sucesso de login e chegada no DET (/servicos)
            for p in reversed(context.pages):
                if not p.is_closed():
                    try:
                        url_p = p.url.lower()
                        if "det.sit.trabalho.gov.br" in url_p or "/servicos" in url_p:
                            card = p.locator(".cardListItem:has(span.title:has-text('CAIXA POSTAL'))")
                            if card.is_visible():
                                page = p
                                sessao_estabelecida = True
                                break
                    except Exception:
                        pass
            if sessao_estabelecida:
                break

            time.sleep(1.0)

        if tela_2fa_detectada:
            registrar_pulo_2fa(
                lista_cnpjs=lista_cnpjs,
                mapa_itens=mapa_itens,
                nome_cert_exibicao=nome_cert_exibicao,
                rodada=rodada,
                num_grupo=num_grupo,
                total_grupos=total_grupos,
                total_itens=total_itens,
                item_global_counter=item_global_counter,
                pasta_cert=pasta_cert,
                base_dir=base_dir,
                execucao_dt_str=execucao_dt_str,
                pagina_2fa=pagina_2fa,
                progress_callback=progress_callback,
            )
            fechar_navegador(browser, context, page)
            if playwright_mgr:
                try:
                    playwright_mgr.stop()
                except Exception:
                    pass
            return False

        if not sessao_estabelecida:
            logger.error(
                f"Tempo limite esgotado ({timeout_login}s) aguardando login ou portal DET para '{nome_cert_exibicao}'."
            )
            for doc in lista_cnpjs:
                it = mapa_itens[doc]
                msg = "Tempo limite esgotado aguardando carregamento do portal DET"
                if rodada == 1:
                    it.status_t1 = "FALHA_SISTEMA"
                    it.detalhe_t1 = msg
                else:
                    it.status_t2 = "FALHA_SISTEMA"
                    it.detalhe_t2 = msg
                it.status_final = "FALHA_SISTEMA"
                it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

            item_global_counter[0] += len(lista_cnpjs)
            if progress_callback:
                progress_callback(
                    item_global_counter[0],
                    total_itens,
                    f"Grupo {num_grupo}/{total_grupos} falhou: Timeout no portal DET",
                )
            fechar_navegador(browser, context, page)
            if playwright_mgr:
                try:
                    playwright_mgr.stop()
                except Exception:
                    pass
            return False

        time.sleep(DELAY_ENTRE_ACOES_S)
        logger.info(f"Fase 2 estabelecida: Sessão ativa para certificado '{nome_cert_exibicao}'!")

        # ==================================================================
        # FASE 3: Loop de CNPJs deste certificado
        # ==================================================================
        perfil_ref = [0]
        falhas_seguidas = 0

        for current_item in lista_cnpjs:
            if stop_event and stop_event.is_set():
                logger.warning("Execução cancelada pelo usuário.")
                break

            if falhas_seguidas >= 3:
                logger.error(
                    f"3 falhas seguidas no grupo '{nome_cert_exibicao}'. Interrompendo itens restantes deste certificado."
                )
                break

            item_id = current_item.strip()
            if not item_id:
                continue

            it = mapa_itens[item_id]
            item_global_counter[0] += 1
            idx_atual = item_global_counter[0]

            if progress_callback:
                progress_callback(
                    idx_atual - 1,
                    total_itens,
                    f"R{rodada}: Processando {idx_atual}/{total_itens} ({item_id}) [Cert: {nome_cert_exibicao[:20]}]",
                )

            # Executa com retry imediato em caso de erro temporário
            sucesso, status_code, detalhe = processar_item_det_com_retry(
                page=page,
                item_id=item_id,
                idx_item=idx_atual,
                total_itens=total_itens,
                screenshots_dir=pasta_cert,
                hoje_date=hoje_date,
                perfil_ref=perfil_ref,
                progress_callback=progress_callback,
                stop_event=stop_event,
            )

            # Registra no ItemExecucao correspondente
            if rodada == 1:
                it.status_t1 = status_code
                it.detalhe_t1 = detalhe
            else:
                it.status_t2 = status_code
                it.detalhe_t2 = detalhe
            it.status_final = status_code
            it.data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

            # Atualiza o relatório geral consolidado com os resultados mais recentes
            try:
                salvar_relatorio_execucao(base_dir, execucao_dt_str, list(mapa_itens.values()))
            except Exception:
                pass

            if sucesso:
                falhas_seguidas = 0
            else:
                falhas_seguidas += 1

            if progress_callback:
                progress_callback(
                    idx_atual,
                    total_itens,
                    f"R{rodada}: Item {idx_atual}/{total_itens} ({item_id}) - {status_code}",
                )

        return True

    except Exception as e_grupo:
        logger.error(f"Erro no processamento do grupo de certificado '{nome_cert_exibicao}': {e_grupo}")
        return False

    finally:
        # FASE 4: Logout do DET e Fechamento Limpo do Chrome deste grupo
        if page and not page.is_closed():
            efetuar_logout_det(page)

        logger.info(f"Fechando sessão do Chrome para certificado '{nome_cert_exibicao}'...")
        fechar_navegador(browser, context, page)
        if playwright_mgr:
            try:
                playwright_mgr.stop()
            except Exception:
                pass
        time.sleep(2)


def executar_robo_hibrido(
    base_dir: Path,
    progress_callback: Optional[Callable[[int, int, str], None]] = None,
    stop_event: Optional[threading.Event] = None,
    planilha_especifica: Optional[str] = None,
):
    if not base_dir:
        raise ValueError("O parâmetro 'base_dir' é obrigatório para executar o robô.")
    base_dir = Path(base_dir).resolve()

    if planilha_especifica:
        p = Path(planilha_especifica)
        excel_path = p if p.is_absolute() else base_dir / p
    else:
        excel_path = base_dir / "CNPJs.xlsx"
        if not excel_path.exists():
            legado_path = base_dir / "CNPJs-e-CPFs.xlsx"
            if legado_path.exists():
                excel_path = legado_path

    hoje_date = datetime.now().date()
    execucao_dt_str = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
    screenshots_raiz = base_dir / "CapturasDeTela-Downloads"
    screenshots_dir = screenshots_raiz / execucao_dt_str
    pasta_logs = base_dir / "Logs"
    pasta_relatorios = base_dir / "Relatorios"

    # Garante a existência prévia das pastas da automação
    pasta_logs.mkdir(parents=True, exist_ok=True)
    screenshots_raiz.mkdir(parents=True, exist_ok=True)
    pasta_relatorios.mkdir(parents=True, exist_ok=True)

    # Rotatividade de histórico: exclui itens com mais de 60 dias se houver 10 ou mais itens
    executar_rotatividade_antigos(pasta_logs, eh_pasta=False, min_itens=10, limite_dias=60)
    executar_rotatividade_antigos(screenshots_raiz, eh_pasta=True, min_itens=10, limite_dias=60)
    executar_rotatividade_antigos(pasta_relatorios, eh_pasta=False, min_itens=10, limite_dias=60)

    # Garante que qualquer resíduo anterior do Chrome seja encerrado antes do início
    liberar_porta_se_ocupada()

    # 1. Verificação da Planilha de Entrada (se não existir, é criada automaticamente se for a padrão)
    if not excel_path.exists():
        if planilha_especifica:
            logger.info(f"A planilha '{excel_path.name}' não foi encontrada em '{base_dir}'. Não há itens pendentes a processar.")
            return True
        logger.error(f"Planilha de CNPJs não encontrada em: {excel_path}")
        try:
            base_dir.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "CNPJs_consultados"
            ws["A1"] = "CNPJ"
            ws["B1"] = "Certificado"
            cor_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            fonte_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            alinhamento_centro = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 24
            ws["A1"].fill = cor_header
            ws["A1"].font = fonte_header
            ws["A1"].alignment = alinhamento_centro
            ws["B1"].fill = cor_header
            ws["B1"].font = fonte_header
            ws["B1"].alignment = alinhamento_centro
            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 65
            wb.save(str(excel_path))
            wb.close()
            logger.info(f"Planilha modelo criada automaticamente com sucesso em: {excel_path}")
        except Exception as e:
            logger.error(f"Falha ao criar planilha modelo: {e}")
        exibir_alerta_planilha_inexistente(excel_path)
        return False

    # Leitura das linhas da planilha: Coluna A = CNPJ, Coluna B = Certificado
    try:
        wb = openpyxl.load_workbook(str(excel_path), data_only=True)
        if "CNPJs_consultados" in wb.sheetnames:
            sheet = wb["CNPJs_consultados"]
        elif "Documentos" in wb.sheetnames:
            sheet = wb["Documentos"]
        else:
            sheet = wb.active
        linhas_lidas = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row:
                continue
            col_a = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            col_b = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if col_a:
                linhas_lidas.append((col_a, col_b))
        wb.close()
    except Exception as e:
        logger.error(f"Erro ao ler a planilha {excel_path}: {e}")
        return False

    # Filtra cabeçalhos conhecidos
    header_names = {
        "CNPJ_OU_CPF", "CNPJ_E_CPF", "CNPJS_E_CPFS", "CNPJ", "CPF",
        "DOCUMENTO", "DOCUMENTOS", "CNPJS_CONSULTADOS"
    }
    dados_filtrados = [
        (doc, cert) for doc, cert in linhas_lidas
        if doc.upper() not in header_names
    ]

    if not dados_filtrados:
        logger.warning(
            f"A planilha em '{excel_path}' não contém CNPJs para processar. Preencha e inicie novamente."
        )
        return False

    # Inicializa o mapa de rastreamento de cada CNPJ para o relatório de execução
    mapa_itens: dict[str, ItemExecucao] = {
        doc: ItemExecucao(doc, cert)
        for doc, cert in dados_filtrados
    }

    # Agrupamento por certificado de todos os CNPJs da planilha (sem validação de formato prévia)
    grupos_certificados = defaultdict(list)
    for doc, cert in dados_filtrados:
        cert_chave = cert.strip()
        grupos_certificados[cert_chave].append(doc)

    total_itens = len(dados_filtrados)
    total_grupos = len(grupos_certificados)

    logger.info(
        f"Planilha carregada com sucesso. Total de CNPJs: {total_itens} "
        f"distribuídos em {total_grupos} grupo(s) de certificado."
    )
    if progress_callback:
        progress_callback(0, total_itens, f"Iniciando automação... (0/{total_itens})")

    limpar_ou_criar_pasta(screenshots_dir)
    logger.info(f"Pasta de capturas desta execução: {screenshots_dir}")

    item_counter = [0]

    try:
        # ==================================================================
        # RODADA 1: Execução Principal de todos os grupos de certificado
        # ==================================================================
        for num_grupo, (cert_alvo, lista_cnpjs) in enumerate(grupos_certificados.items(), start=1):
            if stop_event and stop_event.is_set():
                logger.warning("Execução interrompida pelo usuário.")
                break

            processar_sessao_certificado(
                cert_alvo=cert_alvo,
                lista_cnpjs=lista_cnpjs,
                mapa_itens=mapa_itens,
                rodada=1,
                num_grupo=num_grupo,
                total_grupos=total_grupos,
                total_itens=total_itens,
                item_global_counter=item_counter,
                screenshots_dir=screenshots_dir,
                hoje_date=hoje_date,
                base_dir=base_dir,
                execucao_dt_str=execucao_dt_str,
                progress_callback=progress_callback,
                stop_event=stop_event,
            )

        # ==================================================================
        # RODADA 2: Reprocessamento Automático de CNPJs Pendentes
        # ==================================================================
        itens_pendentes = [it for it in mapa_itens.values() if it.status_final in ("FALHA_SISTEMA", "NAO_EXECUTADO")]
        if itens_pendentes and not (stop_event and stop_event.is_set()):
            logger.info("\n" + "=" * 80)
            logger.info(f"[RODADA 2 - REPROCESSAMENTO AUTOMÁTICO] {len(itens_pendentes)} CNPJ(s) pendente(s).")
            logger.info("Executando segunda passada para recuperar itens com falha temporária...")
            logger.info("=" * 80 + "\n")

            grupos_pendentes = defaultdict(list)
            for it in itens_pendentes:
                grupos_pendentes[it.certificado].append(it.cnpj)

            total_grupos_p = len(grupos_pendentes)
            item_counter_r2 = [0]

            for num_grupo, (cert_alvo, lista_p) in enumerate(grupos_pendentes.items(), start=1):
                if stop_event and stop_event.is_set():
                    break

                processar_sessao_certificado(
                    cert_alvo=cert_alvo,
                    lista_cnpjs=lista_p,
                    mapa_itens=mapa_itens,
                    rodada=2,
                    num_grupo=num_grupo,
                    total_grupos=total_grupos_p,
                    total_itens=len(itens_pendentes),
                    item_global_counter=item_counter_r2,
                    screenshots_dir=screenshots_dir,
                    hoje_date=hoje_date,
                    base_dir=base_dir,
                    execucao_dt_str=execucao_dt_str,
                    progress_callback=progress_callback,
                    stop_event=stop_event,
                )

    finally:
        # Garante o fechamento definitivo de qualquer processo restante do Chrome
        matar_processos_chrome_automacao()
        # Gera o Relatório de Execução consolidado com histórico e a Planilha de Pendentes sempre!
        salvar_relatorio_execucao(base_dir, execucao_dt_str, list(mapa_itens.values()))
        salvar_planilha_pendentes(base_dir, list(mapa_itens.values()))

    itens_sucesso = sum(1 for it in mapa_itens.values() if it.status_final == "SUCESSO")
    itens_2fa = sum(1 for it in mapa_itens.values() if it.status_final == "EXIGE_2FA")
    itens_falha = total_itens - itens_sucesso
    if itens_falha == 0:
        logger.info(f"Automação finalizada com sucesso total: {itens_sucesso}/{total_itens} CNPJ(s) processados com êxito.")
        return True
    else:
        logger.warning(
            f"Automação finalizada com pendências: {itens_sucesso} sucesso(s), "
            f"{itens_2fa} bloqueado(s) por 2FA, "
            f"{itens_falha} pendência(s)/falha(s) de {total_itens} CNPJ(s)."
        )
        return False




if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(message)s",
        datefmt="%d/%m/%Y %H:%M:%S"
    )
    executar_robo_hibrido()

