"""
Módulo de Coleta e Gestão de Certificados Digitais do Windows.
Extrai os certificados instalados no repositório pessoal (Cert:\\CurrentUser\\My)
e mantém a aba 'Certificados_Disponiveis' na planilha Excel (CNPJs.xlsx) atualizada.
"""

import json
import logging
import re
import subprocess
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("RoboSocietario")


def obter_certificados_instalados_windows() -> List[Dict[str, Any]]:
    """
    Consulta o repositório de certificados pessoais do Windows (CurrentUser\\My)
    utilizando PowerShell e retorna lista de certificados formatados.
    """
    ps_cmd = (
        "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8; "
        "Get-ChildItem Cert:\\CurrentUser\\My | "
        "ForEach-Object { [PSCustomObject]@{ "
        "Subject = $_.Subject; "
        "Validade = $_.NotAfter.ToString('dd/MM/yyyy'); "
        "Expirado = ($_.NotAfter -lt (Get-Date)) "
        "} } | ConvertTo-Json"
    )

    try:
        res = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps_cmd],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=15,
        )
    except Exception as e:
        logger.warning(f"[CERTIFICADOS] Erro ao consultar certificados via PowerShell: {e}")
        return []

    if res.returncode != 0 or not res.stdout.strip():
        logger.warning(f"[CERTIFICADOS] PowerShell retornou código {res.returncode} ou saída vazia.")
        return []

    try:
        data = json.loads(res.stdout)
    except Exception as e:
        logger.warning(f"[CERTIFICADOS] Falha ao decodificar JSON dos certificados: {e}")
        return []

    if isinstance(data, dict):
        data = [data]

    certificados = []
    for item in data:
        subj = str(item.get("Subject") or "").strip()
        validade = str(item.get("Validade") or "").strip()
        expirado = bool(item.get("Expirado", False))

        # Extrai CN=...
        cn_match = re.search(r"CN=([^,]+)", subj)
        nome_completo = cn_match.group(1).strip() if cn_match else subj

        # No Brasil, certificados e-CNPJ/e-CPF ICP-Brasil seguem:
        # "RAZAO SOCIAL:CNPJ" ou "NOME PESSOA:CPF"
        if ":" in nome_completo:
            partes = nome_completo.rsplit(":", 1)
            razao_social = partes[0].strip()
            doc = partes[1].strip()
        else:
            razao_social = nome_completo
            doc = ""

        status = "EXPIRADO" if expirado else "VALIDO"
        certificados.append(
            {
                "razao_social": razao_social,
                "documento": doc,
                "nome_completo": nome_completo,
                "validade": validade,
                "status": status,
            }
        )

    return certificados


def formatar_aba_cnpjs_consultados(ws):
    """Aplica o mesmo padrão de design premium na aba CNPJs_consultados."""
    cor_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fonte_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_esq = Alignment(horizontal="left", vertical="center")
    borda_fina = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.row_dimensions[1].height = 24
    for col_idx in [1, 2]:
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = cor_header
        cell.font = fonte_header
        cell.alignment = alinhamento_centro

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        cell_cnpj = ws.cell(row=row_idx, column=1)
        cell_cert = ws.cell(row=row_idx, column=2)

        cell_cnpj.font = Font(name="Calibri", size=11)
        cell_cnpj.alignment = alinhamento_centro
        cell_cnpj.border = borda_fina
        cell_cnpj.number_format = "@"

        cell_cert.font = Font(name="Calibri", size=11)
        cell_cert.alignment = alinhamento_esq
        cell_cert.border = borda_fina

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 65


def atualizar_aba_certificados_disponiveis(caminho_planilha: Path) -> int:
    """
    Atualiza ou cria a aba 'Certificados_Disponiveis' na planilha Excel informada,
    mantendo a aba principal ('CNPJs_consultados') intacta e formatada no mesmo estilo.

    Retorna a quantidade de certificados gravados.
    """
    caminho_planilha = Path(caminho_planilha).resolve()
    certificados = obter_certificados_instalados_windows()

    if caminho_planilha.exists():
        try:
            wb = openpyxl.load_workbook(str(caminho_planilha))
        except PermissionError:
            raise PermissionError(
                f"A planilha '{caminho_planilha.name}' está aberta no Excel. "
                "Feche o arquivo antes de atualizar os certificados."
            )
        # Se ainda possuir a aba com nome legado 'Documentos', renomeia para 'CNPJs_consultados'
        if "Documentos" in wb.sheetnames and "CNPJs_consultados" not in wb.sheetnames:
            wb["Documentos"].title = "CNPJs_consultados"
    else:
        wb = openpyxl.Workbook()
        ws_doc = wb.active
        ws_doc.title = "CNPJs_consultados"
        ws_doc.append(["CNPJ", "Certificado"])

    # Formata a aba principal CNPJs_consultados se ela existir
    if "CNPJs_consultados" in wb.sheetnames:
        formatar_aba_cnpjs_consultados(wb["CNPJs_consultados"])

    nome_aba = "Certificados_Disponiveis"
    if nome_aba in wb.sheetnames:
        del wb[nome_aba]

    ws = wb.create_sheet(title=nome_aba)

    # Cabeçalhos
    headers = [
        "Razão Social / Nome",
        "CNPJ / CPF do Certificado",
        "Nome Completo (certmgr)",
        "Validade",
        "Status",
    ]
    ws.append(headers)

    # Linhas de certificados
    for c in certificados:
        ws.append(
            [
                c["razao_social"],
                c["documento"],
                c["nome_completo"],
                c["validade"],
                c["status"],
            ]
        )

    # Estilização visual dos cabeçalhos e células
    cor_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fonte_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_esq = Alignment(horizontal="left", vertical="center")
    borda_fina = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Estiliza cabeçalho
    ws.row_dimensions[1].height = 24
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = cor_header
        cell.font = fonte_header
        cell.alignment = alinhamento_centro

    # Estiliza conteúdo
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)),
        start=2,
    ):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, cell in enumerate(row, start=1):
            cell.border = borda_fina
            if col_idx in [2, 4, 5]:
                cell.alignment = alinhamento_centro
            else:
                cell.alignment = alinhamento_esq

            if col_idx == 5:
                if cell.value == "VALIDO":
                    cell.font = Font(name="Calibri", bold=True, color="006100")
                    cell.fill = PatternFill(
                        start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
                    )
                else:
                    cell.font = Font(name="Calibri", bold=True, color="9C0006")
                    cell.fill = PatternFill(
                        start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
                    )

    # Ajusta largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # Deixa a aba 'CNPJs_consultados' como a ativa ao abrir o Excel
    if "CNPJs_consultados" in wb.sheetnames:
        wb.active = wb["CNPJs_consultados"]
    elif "Documentos" in wb.sheetnames:
        wb.active = wb["Documentos"]

    try:
        wb.save(str(caminho_planilha))
        wb.close()
    except PermissionError:
        raise PermissionError(
            f"A planilha '{caminho_planilha.name}' está aberta no Excel. "
            "Feche o arquivo antes de atualizar os certificados."
        )

    return len(certificados)
