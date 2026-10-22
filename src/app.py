import sys
import warnings

# Define STA (Single-Threaded Apartment) antes de qualquer importação de COM/pywinauto.
# Isso impede conflitos que travam diálogos nativos do Windows (como filedialog.askdirectory).
if not hasattr(sys, "coinit_flags"):
    sys.coinit_flags = 2
warnings.filterwarnings("ignore", message=".*Apply externally defined coinit_flags.*")

import os
import threading
import logging
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk
import openpyxl

try:
    from src.automacao_hibrida_final import BASE_DIR as BASE_DIR_PADRAO, executar_robo_hibrido
except ImportError:
    from automacao_hibrida_final import BASE_DIR as BASE_DIR_PADRAO, executar_robo_hibrido

# Configuração global de aparência do CustomTkinter
ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")


class TextboxLogHandler(logging.Handler):
    """
    Handler thread-safe para redirecionar logs diretamente
    para o CTkTextbox da interface gráfica.
    """
    def __init__(self, textbox: ctk.CTkTextbox, root: ctk.CTk):
        super().__init__()
        self.textbox = textbox
        self.root = root

    def emit(self, record: logging.LogRecord):
        try:
            msg = self.format(record)
            self.root.after(0, self._append_msg, msg)
        except Exception:
            pass

    def _append_msg(self, msg: str):
        try:
            self.textbox.configure(state="normal")
            self.textbox.insert("end", msg + "\n")
            self.textbox.see("end")
            self.textbox.configure(state="disabled")
        except Exception:
            pass


class AppDET(ctk.CTk):
    def __init__(self):
        super().__init__()

        # Configurações da Janela Principal
        self.title("Automação DET - Consulta de Notificações e Certificados")
        self.geometry("620x315")
        self.minsize(560, 280)

        # Variáveis de Controle
        self.caminho_var = ctk.StringVar(value=str(BASE_DIR_PADRAO))
        self.stop_event = threading.Event()
        self.worker_thread = None
        self.executando = False

        # Configuração do Logger
        self.logger = logging.getLogger("RoboDET")
        self.logger.setLevel(logging.INFO)

        # Construção da Interface
        self._criar_layout()
        self._configurar_logger()

    def _criar_layout(self):
        # Container principal com padding compacto
        self.main_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.main_frame.pack(fill="both", expand=True, padx=12, pady=10)

        # ======================================================================
        # 1. LINHA SUPERIOR: Caminho, Procurar, Iniciar e Cancelar
        # ======================================================================
        self.top_frame = ctk.CTkFrame(self.main_frame, fg_color="#f8fafc", corner_radius=8, border_width=1, border_color="#e2e8f0")
        self.top_frame.pack(fill="x", pady=(0, 6), ipady=3, ipadx=4)

        # Input do caminho
        self.entry_caminho = ctk.CTkEntry(
            self.top_frame,
            textvariable=self.caminho_var,
            font=("Segoe UI", 11),
            height=30,
            placeholder_text="Selecione o diretório onde a pasta deve ser gerada...",
            fg_color="#ffffff",
            text_color="#0f172a",
            border_color="#cbd5e1"
        )
        self.entry_caminho.pack(side="left", fill="x", expand=True, padx=(6, 6))

        # Botão Procurar (Azul #0284c7)
        self.btn_procurar = ctk.CTkButton(
            self.top_frame,
            text="Procurar ▾",
            width=85,
            height=30,
            font=("Segoe UI", 11, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            text_color="#ffffff",
            command=self._acao_procurar
        )
        self.btn_procurar.pack(side="left", padx=(0, 6))

        # Botão Iniciar Automação (Mesmo azul de procurar: #0284c7)
        self.btn_iniciar = ctk.CTkButton(
            self.top_frame,
            text="Iniciar Automação",
            width=120,
            height=30,
            font=("Segoe UI", 11, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            text_color="#ffffff",
            command=self._acao_iniciar
        )
        self.btn_iniciar.pack(side="left", padx=(0, 6))

        # Botão Cancelar
        self.btn_cancelar = ctk.CTkButton(
            self.top_frame,
            text="Cancelar",
            width=70,
            height=30,
            font=("Segoe UI", 11, "bold"),
            fg_color="#dc2626",
            hover_color="#b91c1c",
            text_color="#ffffff",
            state="disabled",
            command=self._acao_cancelar
        )
        self.btn_cancelar.pack(side="left", padx=(0, 6))

        # ======================================================================
        # 2. SEÇÃO DE PROGRESSO: Status e Barra
        # ======================================================================
        self.progress_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        self.progress_frame.pack(fill="x", pady=(0, 6))

        # Labels de status e contagem
        self.info_progresso_frame = ctk.CTkFrame(self.progress_frame, fg_color="transparent")
        self.info_progresso_frame.pack(fill="x", pady=(0, 2))

        self.lbl_status = ctk.CTkLabel(
            self.info_progresso_frame,
            text="Pronto para iniciar",
            font=("Segoe UI", 11),
            text_color="#475569"
        )
        self.lbl_status.pack(side="left")

        self.lbl_contador = ctk.CTkLabel(
            self.info_progresso_frame,
            text="0/0 itens • 0%",
            font=("Segoe UI", 11, "bold"),
            text_color="#334155"
        )
        self.lbl_contador.pack(side="right")

        # Barra de Progresso
        self.barra_progresso = ctk.CTkProgressBar(
            self.progress_frame,
            height=8,
            corner_radius=4,
            progress_color="#0284c7",
            fg_color="#e2e8f0"
        )
        self.barra_progresso.pack(fill="x")
        self.barra_progresso.set(0.0)

        # ======================================================================
        # 4. RODAPÉ ANCORADO NA BASE: Botões 'Acessar resultado' e 'Configurar planilha'
        # Ancorado com side="bottom" para NUNCA ser empurrado para fora da janela
        # ======================================================================
        self.bottom_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        self.bottom_frame.pack(side="bottom", fill="x", pady=(4, 0))

        # Configurar Planilha (Verde escuro #166534)
        self.btn_config_planilha = ctk.CTkButton(
            self.bottom_frame,
            text="Configurar planilha",
            height=30,
            font=("Segoe UI", 11, "bold"),
            fg_color="#166534",
            hover_color="#14532d",
            text_color="#ffffff",
            command=self._acao_configurar_planilha
        )
        self.btn_config_planilha.pack(side="right")

        # Acessar Resultado (Ao lado esquerdo de Configurar Planilha)
        self.btn_acessar_resultado = ctk.CTkButton(
            self.bottom_frame,
            text="Acessar resultado",
            height=30,
            font=("Segoe UI", 11, "bold"),
            fg_color="#475569",
            hover_color="#334155",
            text_color="#ffffff",
            command=self._acao_acessar_resultado
        )
        self.btn_acessar_resultado.pack(side="right", padx=(0, 8))

        # ======================================================================
        # 3. SEÇÃO CENTRAL: Log de Execução (Ocupa o espaço intermediário)
        # ======================================================================
        self.log_container = ctk.CTkFrame(self.main_frame, fg_color="#ffffff", corner_radius=8, border_width=1, border_color="#e2e8f0")
        self.log_container.pack(fill="both", expand=True, pady=(0, 6))

        self.lbl_titulo_log = ctk.CTkLabel(
            self.log_container,
            text="Log de execução",
            font=("Segoe UI", 11, "bold"),
            text_color="#1e293b"
        )
        self.lbl_titulo_log.pack(anchor="w", padx=8, pady=(4, 2))

        self.txt_log = ctk.CTkTextbox(
            self.log_container,
            height=80,
            font=("Consolas", 10),
            fg_color="#f8fafc",
            text_color="#0f172a",
            corner_radius=6,
            border_width=1,
            border_color="#cbd5e1",
            wrap="none"
        )
        self.txt_log.pack(fill="both", expand=True, padx=8, pady=(0, 6))
        self.txt_log.configure(state="disabled")

    def _configurar_logger(self):
        """Redireciona os registros do logger RoboDET para a caixa de texto da UI."""
        for h in list(self.logger.handlers):
            self.logger.removeHandler(h)

        handler = TextboxLogHandler(self.txt_log, self)
        formatter = logging.Formatter("%(asctime)s - %(message)s", datefmt="%d/%m/%Y %H:%M:%S")
        handler.setFormatter(formatter)
        self.logger.addHandler(handler)

    def _acao_procurar(self):
        """Abre o seletor de diretório do Windows com suporte modal correto e foco garantido."""
        caminho_inicial = self.caminho_var.get().strip()
        try:
            if not caminho_inicial or not Path(caminho_inicial).exists():
                caminho_inicial = str(Path.home())
        except Exception:
            caminho_inicial = str(Path.home())

        pasta_selecionada = filedialog.askdirectory(
            parent=self,
            initialdir=caminho_inicial,
            title="Selecione o diretório de destino da automação"
        )
        if pasta_selecionada:
            caminho_formatado = str(Path(pasta_selecionada))
            self.caminho_var.set(caminho_formatado)
            self.logger.info(f"Diretório selecionado: {caminho_formatado}")

        self.lift()
        self.focus_force()

    def _acao_acessar_resultado(self):
        """Abre a pasta 'CapturasDeTela' no Windows Explorer."""
        caminho_str = self.caminho_var.get().strip()
        pasta_base = Path(caminho_str) if caminho_str else BASE_DIR_PADRAO
        pasta_prints = pasta_base / "CapturasDeTela"
        try:
            pasta_prints.mkdir(parents=True, exist_ok=True)
            self.logger.info(f"Abrindo pasta de resultados: {pasta_prints}")
            os.startfile(str(pasta_prints))
        except Exception as e:
            self.logger.error(f"Erro ao acessar pasta de capturas: {e}")
            messagebox.showerror("Erro", f"Não foi possível abrir a pasta de resultados:\n{e}", parent=self)

    def _acao_configurar_planilha(self):
        """
        Abre o arquivo Excel 'CNPJs.xlsx' no diretório selecionado.
        Caso não exista, cria com o cabeçalho 'CNPJ', avisa no log e abre.
        """
        caminho_str = self.caminho_var.get().strip()
        if not caminho_str:
            messagebox.showerror("Aviso", "Por favor, especifique um diretório válido.", parent=self)
            return

        pasta_base = Path(caminho_str)
        try:
            pasta_base.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            messagebox.showerror("Erro de Permissão", f"Não foi possível acessar ou criar a pasta:\n{e}", parent=self)
            return

        caminho_excel = pasta_base / "CNPJs.xlsx"
        if not caminho_excel.exists():
            legado_excel = pasta_base / "CNPJs-e-CPFs.xlsx"
            if legado_excel.exists():
                caminho_excel = legado_excel

        if not caminho_excel.exists():
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Documentos"
                ws["A1"] = "CNPJ"
                ws["B1"] = "Certificado"
                wb.save(str(caminho_excel))
                wb.close()
                self.logger.info(f"Planilha 'CNPJs.xlsx' criada com sucesso em: {caminho_excel}")
            except Exception as e:
                self.logger.error(f"Falha ao criar planilha: {e}")
                messagebox.showerror("Erro", f"Não foi possível criar o arquivo Excel:\n{e}", parent=self)
                return
        else:
            self.logger.info(f"Abrindo planilha existente em: {caminho_excel}")

        try:
            os.startfile(str(caminho_excel))
        except Exception as e:
            self.logger.error(f"Erro ao abrir o Excel: {e}")
            messagebox.showerror("Erro", f"Não foi possível iniciar o Excel automaticamente:\n{e}", parent=self)

    def _acao_iniciar(self):
        """Inicia a execução da automação em uma thread paralela."""
        caminho_str = self.caminho_var.get().strip()
        if not caminho_str:
            messagebox.showwarning("Aviso", "Informe um caminho válido para a geração da pasta.", parent=self)
            return

        pasta_base = Path(caminho_str)
        try:
            pasta_base.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            messagebox.showerror("Erro de Pasta", f"Não foi possível acessar a pasta selecionada:\n{e}", parent=self)
            return

        caminho_excel = pasta_base / "CNPJs.xlsx"
        if not caminho_excel.exists() and (pasta_base / "CNPJs-e-CPFs.xlsx").exists():
            caminho_excel = pasta_base / "CNPJs-e-CPFs.xlsx"

        if not caminho_excel.exists():
            resp = messagebox.askyesno(
                "Planilha não encontrada",
                f"A planilha 'CNPJs.xlsx' não existe na pasta selecionada:\n{caminho_excel}\n\n"
                "Deseja criá-la e abri-la agora para preencher?",
                parent=self
            )
            if resp:
                self._acao_configurar_planilha()
            return

        # Ajusta estado dos botões
        self.executando = True
        self.stop_event.clear()
        self.btn_iniciar.configure(state="disabled")
        self.btn_procurar.configure(state="disabled")
        self.btn_config_planilha.configure(state="disabled")
        self.btn_acessar_resultado.configure(state="disabled")
        self.entry_caminho.configure(state="disabled")
        self.btn_cancelar.configure(state="normal", text="Cancelar")

        self.barra_progresso.set(0.0)
        self.lbl_contador.configure(text="0/0 itens • 0%")
        self.lbl_status.configure(text="Iniciando execução...")

        # Dispara thread de automação
        self.worker_thread = threading.Thread(
            target=self._executar_worker,
            args=(pasta_base,),
            daemon=True
        )
        self.worker_thread.start()

    def _acao_cancelar(self):
        """Sinaliza parada para a thread de automação."""
        if self.executando:
            self.stop_event.set()
            self.btn_cancelar.configure(state="disabled", text="Cancelando...")
            self.lbl_status.configure(text="Cancelando automação...")
            self.logger.warning("Solicitação de cancelamento enviada. Finalizando rotina com segurança...")

    def _atualizar_progresso(self, atual: int, total: int, status_texto: str):
        """Callback thread-safe para atualizar a barra de progresso e textos."""
        def _sync():
            frac = (atual / total) if total > 0 else 0.0
            frac = max(0.0, min(1.0, frac))
            self.barra_progresso.set(frac)
            porcentagem = int(frac * 100)
            self.lbl_contador.configure(text=f"{atual}/{total} itens • {porcentagem}%")
            self.lbl_status.configure(text=status_texto)

        self.after(0, _sync)

    def _executar_worker(self, pasta_base: Path):
        """Executa a rotina do robô em background."""
        try:
            executar_robo_hibrido(
                base_dir=pasta_base,
                progress_callback=self._atualizar_progresso,
                stop_event=self.stop_event
            )
        except Exception as e:
            self.logger.error(f"Erro inesperado durante a execução: {e}")
        finally:
            self.after(0, self._finalizar_execucao)

    def _finalizar_execucao(self):
        """Restaura a interface após o encerramento do robô."""
        self.executando = False
        self.btn_iniciar.configure(state="normal")
        self.btn_procurar.configure(state="normal")
        self.btn_config_planilha.configure(state="normal")
        self.btn_acessar_resultado.configure(state="normal")
        self.entry_caminho.configure(state="normal")
        self.btn_cancelar.configure(state="disabled", text="Cancelar")

        if self.stop_event.is_set():
            self.lbl_status.configure(text="Execução cancelada.")
        else:
            self.lbl_status.configure(text="Execução finalizada.")


def main():
    app = AppDET()
    app.mainloop()


if __name__ == "__main__":
    main()
